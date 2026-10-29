"""
Habeas Corpus — modules/clientes.py
Clientes del despacho: lista con lo que importa (asuntos en curso, saldo, último
contacto), alta con validación de RFC/CURP y aviso de duplicados, ficha con bitácora,
asuntos, cobranza, documentos y datos; importación desde Excel y ficha en PDF.
"""

from __future__ import annotations

import html
import io
import json
import math
import re
from datetime import timedelta
from urllib.parse import quote

from flask import (
    Blueprint, abort, g, jsonify, redirect, render_template,
    request, send_file, session, url_for,
)

from modules import boveda, cobranzas, lector
from modules.asuntos import (
    ABIERTOS, ESTADOS as ESTADOS_ASUNTO, TIPOS_DERECHO, _FILTRO_PROPIOS,
    _fecha, _hoy, _local, _solo_lo_suyo, _usuarios,
)
from modules.auth import login_required, role_required, tenant_required

clientes_bp = Blueprint("clientes", __name__)

# ─────────────────────────────────────────────────────────────────────────────
#  Constantes
# ─────────────────────────────────────────────────────────────────────────────

ESTADOS_MX = [
    "Aguascalientes", "Baja California", "Baja California Sur", "Campeche",
    "Chiapas", "Chihuahua", "Ciudad de México", "Coahuila de Zaragoza",
    "Colima", "Durango", "Estado de México", "Guanajuato", "Guerrero",
    "Hidalgo", "Jalisco", "Michoacán de Ocampo", "Morelos", "Nayarit",
    "Nuevo León", "Oaxaca", "Puebla", "Querétaro", "Quintana Roo",
    "San Luis Potosí", "Sinaloa", "Sonora", "Tabasco", "Tamaulipas",
    "Tlaxcala", "Veracruz de Ignacio de la Llave", "Yucatán", "Zacatecas",
]

ORIGENES = [
    "Referido por cliente", "Referido por colega", "Búsqueda web",
    "Redes sociales", "Publicidad digital", "Directorio legal",
    "Evento / conferencia", "Otro",
]

TIPOS = {"persona_fisica": "Persona física", "persona_moral": "Empresa"}

ESTATUS = {
    "activo":    {"label": "Cliente",   "color": "#3D5A80", "fondo": "#EEF1F5"},
    "prospecto": {"label": "Prospecto", "color": "#8A6A2F", "fondo": "#F7F0E1"},
    "inactivo":  {"label": "Inactivo",  "color": "#60718D", "fondo": "#F1EEE8"},
}

GRUPOS = {"activo": "Clientes", "prospecto": "Prospectos", "inactivo": "Inactivos", "todos": "Todos"}

BITACORA = {
    "nota":     "Nota",
    "llamada":  "Llamada",
    "reunion":  "Reunión",
    "whatsapp": "WhatsApp",
    "correo":   "Correo",
}
TIPOS_NOTA = list(BITACORA) + ["otro"]
TIPOS_COMUNICACION = ["llamada", "correo", "reunion", "whatsapp", "otro"]

NOMBRE_SQL = "COALESCE(NULLIF(c.razon_social, ''), TRIM(c.nombre || ' ' || COALESCE(c.apellidos, '')))"
RE_RFC = re.compile(r"^[A-ZÑ&]{3,4}\d{6}[A-Z\d]{3}$")
RE_CURP = re.compile(r"^[A-Z][AEIOUX][A-Z]{2}\d{6}[HM][A-Z]{5}[A-Z\d]\d$")
RE_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
PER_PAGE = 25

# ─────────────────────────────────────────────────────────────────────────────
#  Helpers
# ─────────────────────────────────────────────────────────────────────────────


def _log(accion: str, entidad: str = "cliente", entidad_id: int | None = None, detalles: dict | None = None):
    try:
        g.tenant_db.execute(
            "INSERT INTO actividad_log (user_id, accion, entidad, entidad_id, detalles, ip) VALUES (?,?,?,?,?,?)",
            (session.get("user_id"), accion, entidad, entidad_id,
             json.dumps(detalles or {}, ensure_ascii=False), request.remote_addr),
        )
        g.tenant_db.commit()
    except Exception:
        pass


def _safe_int(v, default=0):
    try:
        return int(v)
    except (TypeError, ValueError):
        return default


def _safe_float(v, default=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def _get_abogados():
    try:
        rows = g.global_db.execute(
            "SELECT id, nombre, apellidos FROM tenant_users WHERE tenant_id=? "
            "AND rol IN ('abogado','admin','super_admin') AND activo=1 ORDER BY nombre",
            (g.tenant["id"],),
        ).fetchall()
        return [dict(r) for r in rows]
    except Exception:
        return []


def _es_admin() -> bool:
    return session.get("role") in ("super_admin", "admin")


def nombre_de(c) -> str:
    c = dict(c)
    if c.get("tipo") == "persona_moral" and c.get("razon_social"):
        return c["razon_social"].strip()
    return f"{c.get('nombre') or ''} {c.get('apellidos') or ''}".strip() or "Sin nombre"


def _iniciales(nombre: str) -> str:
    partes = [p for p in re.split(r"\s+", nombre) if p and p[0].isalpha()]
    return ((partes[0][0] + partes[1][0]) if len(partes) > 1 else nombre[:2]).upper()


def _color(cid: int) -> str:
    """Avatar del cliente en la familia marino, gris azulado y dorado."""
    paleta = ("#142E59", "#AD9060", "#34496B", "#60718D", "#8A6A2F", "#3D5A80", "#6E5A3A", "#4A5A73")
    return paleta[cid % len(paleta)]


def _fecha_txt(v, anio=True) -> str:
    f = _fecha(v)
    if not f:
        return "—"
    return f"{f.day} {MESES[f.month - 1]}" + (f" {f.year}" if anio else "")


def _hace(v) -> str | None:
    f = _fecha(v)
    if not f:
        return None
    dias = (_hoy() - f).days
    if dias <= 0:
        return "Hoy"
    if dias == 1:
        return "Ayer"
    if dias < 30:
        return f"Hace {dias} días"
    if dias < 365:
        meses = dias // 30
        return f"Hace {meses} mes{'es' if meses > 1 else ''}"
    return f"Hace {dias // 365} año{'s' if dias >= 730 else ''}"


def _alc_asuntos(alias="a") -> tuple[str, list]:
    """Filtro de asuntos visibles para el usuario (los abogados solo ven los suyos)."""
    if not _solo_lo_suyo():
        return "1", []
    uid = session.get("user_id")
    return _FILTRO_PROPIOS.replace("a.", f"{alias}."), [uid, uid, uid]


def _whatsapp(c, texto: str | None = None) -> str | None:
    tel = re.sub(r"\D", "", c.get("celular") or c.get("telefono") or "")
    if len(tel) < 10:
        return None
    tel = tel if len(tel) > 10 else "52" + tel
    return f"https://wa.me/{tel}" + (f"?text={quote(texto)}" if texto else "")


def _etiquetas(v) -> list:
    try:
        lista = json.loads(v or "[]")
        return [str(x).strip()[:30] for x in lista if str(x).strip()][:12]
    except Exception:
        return []


# ─────────────────────────────────────────────────────────────────────────────
#  LISTA
# ─────────────────────────────────────────────────────────────────────────────


@clientes_bp.route("/")
@login_required
@tenant_required
def index():
    tdb = g.tenant_db
    cobranzas.al_dia(tdb)
    hoy = _hoy()
    a = request.args
    grupo = a.get("estatus") if a.get("estatus") in GRUPOS else "activo"
    q = a.get("q", "").strip()[:100]
    if q and "estatus" not in a:
        grupo = "todos"
    tipo = a.get("tipo") if a.get("tipo") in TIPOS else ""
    abogado = _safe_int(a.get("abogado"), None)
    orden = a.get("orden") if a.get("orden") in ("recientes", "nombre", "saldo", "contacto") else "recientes"

    where, params = ["c.activo=1"], []
    if grupo != "todos":
        where.append("COALESCE(c.estatus,'activo')=?")
        params.append(grupo)
    # Buscador inteligente: sin acentos, tolera errores de dedo y entiende «de Luis», «amparo», etc.
    from modules import buscador
    chips, orden_busqueda, sugerencia = [], None, None
    if q:
        inter = buscador.interpretar(q, _usuarios())
        hallados = buscador.personas(tdb, inter, limite=None)
        chips, orden_busqueda = inter["chips"], hallados["ids"]
        sugerencia = " ".join(hallados["corregido"]) if hallados.get("corregido") else None
        where.append(f"c.id IN ({','.join('?' * len(orden_busqueda)) or 'NULL'})")
        params += orden_busqueda
    if tipo:
        where.append("c.tipo=?")
        params.append(tipo)
    if abogado:
        where.append("c.abogado_id=?")
        params.append(abogado)
    clause = " AND ".join(where)

    alc_a, pa = _alc_asuntos("x")
    alc_c, pc = cobranzas._alcance()
    abiertos = ",".join(f"'{e}'" for e in ABIERTOS)
    columnas = f"""
        SELECT c.*, {NOMBRE_SQL} AS nombre_v,
          (SELECT COUNT(*) FROM asuntos x WHERE x.cliente_id=c.id AND x.estado IN ({abiertos}) AND {alc_a}) AS en_curso,
          (SELECT COUNT(*) FROM asuntos x WHERE x.cliente_id=c.id AND {alc_a}) AS asuntos_total,
          (SELECT COALESCE(SUM({cobranzas.SALDO}),0) FROM cobros co WHERE co.cliente_id=c.id AND {cobranzas.ABIERTO_SQL} AND {alc_c}) AS saldo,
          (SELECT COALESCE(SUM({cobranzas.SALDO}),0) FROM cobros co WHERE co.cliente_id=c.id AND {cobranzas.ABIERTO_SQL}
                AND date(co.fecha_vencimiento) < ? AND {alc_c}) AS vencido,
          MAX(COALESCE((SELECT MAX(fecha) FROM cliente_notas n WHERE n.cliente_id=c.id), ''),
              COALESCE((SELECT MAX(fecha) FROM cliente_comunicaciones m WHERE m.cliente_id=c.id), '')) AS ultimo_contacto
        FROM clientes c WHERE {clause}"""
    col_params = pa + pa + pc + [hoy.isoformat()] + pc
    orden_sql = {"nombre": "nombre_v COLLATE NOCASE", "saldo": "saldo DESC, nombre_v COLLATE NOCASE",
                 "contacto": "ultimo_contacto DESC", "recientes": "c.fecha_alta DESC, c.id DESC"}[orden]

    if orden_busqueda:
        orden_sql = "CASE c.id " + " ".join(f"WHEN {int(i)} THEN {n}" for n, i in enumerate(orden_busqueda)) + " END"
    total = tdb.execute(f"SELECT COUNT(*) FROM clientes c WHERE {clause}", params).fetchone()[0]
    total_pages = max(1, math.ceil(total / PER_PAGE))
    page = min(max(1, _safe_int(a.get("page"), 1)), total_pages)
    usuarios = _usuarios()
    clientes = []
    for r in tdb.execute(f"{columnas} ORDER BY {orden_sql} LIMIT ? OFFSET ?",
                         col_params + params + [PER_PAGE, (page - 1) * PER_PAGE]).fetchall():
        c = dict(r)
        c["iniciales"], c["color"] = _iniciales(c["nombre_v"]), _color(c["id"])
        c["estatus_info"] = ESTATUS.get(c.get("estatus") or "activo", ESTATUS["activo"])
        c["contacto_hace"] = _hace(_local(c["ultimo_contacto"])) if c["ultimo_contacto"] else None
        c["responsable"] = usuarios.get(c.get("abogado_id"))
        c["tel"] = c.get("celular") or c.get("telefono")
        clientes.append(c)

    # Resumen
    conteos = {k: 0 for k in GRUPOS}
    for est, n in tdb.execute("SELECT COALESCE(estatus,'activo'), COUNT(*) FROM clientes WHERE activo=1 GROUP BY 1").fetchall():
        conteos[est] = conteos.get(est, 0) + n
        conteos["todos"] += n
    alc_a2, pa2 = _alc_asuntos("a")
    con_asuntos = tdb.execute(
        f"SELECT COUNT(DISTINCT a.cliente_id) FROM asuntos a JOIN clientes c ON c.id=a.cliente_id "
        f"WHERE c.activo=1 AND a.estado IN ({abiertos}) AND {alc_a2}", pa2).fetchone()[0]
    deuda = tdb.execute(
        f"SELECT COALESCE(SUM({cobranzas.SALDO}),0), COUNT(DISTINCT co.cliente_id) FROM cobros co "
        f"WHERE {cobranzas.ABIERTO_SQL} AND {alc_c}", pc).fetchone()
    nuevos = tdb.execute("SELECT COUNT(*) FROM clientes WHERE activo=1 AND date(fecha_alta) >= ?",
                         (hoy.replace(day=1).isoformat(),)).fetchone()[0]
    resumen = {"con_asuntos": con_asuntos, "saldo": deuda[0], "con_saldo": deuda[1], "nuevos": nuevos}

    base_args = {k: v for k, v in a.items() if k not in ("page", "tenant")}

    def liga(**cambios):
        args = {**base_args, **cambios}
        return url_for("clientes.index", **{k: v for k, v in args.items() if v not in (None, "")})

    return render_template(
        "clientes/index.html",
        clientes=clientes, total=total, page=page, total_pages=total_pages,
        grupo=grupo, q=q, tipo=tipo, abogado=abogado, orden=orden, liga=liga,
        grupos=GRUPOS, conteos=conteos, resumen=resumen, tipos=TIPOS,
        abogados=_get_abogados(), estados_mx=ESTADOS_MX, origenes=ORIGENES,
        es_admin=_es_admin(), volver=a.get("volver", ""), chips=chips, sugerencia=sugerencia,
    )


@clientes_bp.route("/nuevo")
@login_required
@tenant_required
def nuevo():
    return redirect(url_for("clientes.index") + "#nuevo")


# ─────────────────────────────────────────────────────────────────────────────
#  ALTA Y EDICIÓN
# ─────────────────────────────────────────────────────────────────────────────

_CAMPOS_TEXTO = {
    "apellidos": 120, "email": 160, "telefono": 30, "celular": 30, "direccion": 200, "num_ext": 20,
    "num_int": 20, "colonia": 120, "municipio": 120, "ciudad": 120, "cp": 5, "giro_comercial": 120,
    "representante_legal": 160, "contacto_principal": 160, "notas": 3000,
}


def _leer_cliente(tdb, cid: int | None = None):
    """Valida el formulario del cliente. Regresa (valores, error, posibles_duplicados)."""
    d = request.form
    tipo = d.get("tipo") if d.get("tipo") in TIPOS else "persona_fisica"
    v = {k: ((d.get(k) or "").strip()[:n] or None) for k, n in _CAMPOS_TEXTO.items()}
    v["tipo"] = tipo
    v["rfc"] = re.sub(r"[\s-]", "", (d.get("rfc") or "").upper()) or None
    v["curp"] = re.sub(r"\s", "", (d.get("curp") or "").upper()) or None
    if tipo == "persona_moral":
        v["razon_social"] = (d.get("razon_social") or "").strip()[:200] or None
        if not v["razon_social"]:
            return None, "Escribe la razón social de la empresa", []
        v["nombre"], v["apellidos"], v["curp"] = v["razon_social"], None, None
    else:
        v["nombre"] = (d.get("nombre") or "").strip()[:120]
        v["razon_social"] = None
        if not v["nombre"]:
            return None, "Escribe el nombre del cliente", []
    if v["rfc"] and not RE_RFC.match(v["rfc"]):
        return None, "El RFC no tiene un formato válido (12 caracteres para empresa, 13 para persona física)", []
    if v["rfc"] and len(v["rfc"]) != (12 if tipo == "persona_moral" else 13):
        return None, f"El RFC de {'una empresa' if tipo == 'persona_moral' else 'una persona física'} lleva {12 if tipo == 'persona_moral' else 13} caracteres", []
    if v["curp"] and not RE_CURP.match(v["curp"]):
        return None, "La CURP no tiene un formato válido (18 caracteres)", []
    if v["email"]:
        v["email"] = v["email"].lower()
        if not RE_EMAIL.match(v["email"]):
            return None, "El correo no parece válido", []
    if v["cp"] and not re.fullmatch(r"\d{5}", v["cp"]):
        return None, "El código postal lleva 5 dígitos", []
    for campo in ("celular", "telefono"):
        if v[campo] and len(re.sub(r"\D", "", v[campo])) < 10:
            return None, f"El {'celular' if campo == 'celular' else 'teléfono'} debe tener al menos 10 dígitos", []
    v["estado"] = d.get("estado") if d.get("estado") in ESTADOS_MX else None
    v["origen_cliente"] = d.get("origen_cliente") if d.get("origen_cliente") in ORIGENES else None
    v["fecha_nacimiento"] = (_fecha(d.get("fecha_nacimiento")) or None) and d.get("fecha_nacimiento")[:10]
    v["estatus"] = d.get("estatus") if d.get("estatus") in ESTATUS else "activo"
    abogado = _safe_int(d.get("abogado_id"), None)
    v["abogado_id"] = abogado if abogado in {a["id"] for a in _get_abogados()} else None
    v["etiquetas"] = json.dumps([t.strip()[:30] for t in (d.get("etiquetas") or "").split(",") if t.strip()][:12],
                                ensure_ascii=False)

    # Posibles duplicados: mismo RFC, correo o celular
    dup, cond, p = [], [], []
    if v["rfc"]:
        cond.append("rfc=?"); p.append(v["rfc"])
    if v["email"]:
        cond.append("LOWER(email)=?"); p.append(v["email"])
    tel = re.sub(r"\D", "", v["celular"] or "")[-10:]
    if len(tel) == 10:
        cond.append("SUBSTR(REPLACE(REPLACE(REPLACE(celular,' ',''),'-',''),'+',''), -10)=?"); p.append(tel)
    if cond:
        for r in tdb.execute(f"SELECT c.*, {NOMBRE_SQL} AS nombre_v FROM clientes c WHERE activo=1 AND id != ? "
                             f"AND ({' OR '.join(cond)}) LIMIT 3", (cid or 0, *p)).fetchall():
            motivo = ("el mismo RFC" if v["rfc"] and r["rfc"] == v["rfc"] else
                      "el mismo correo" if v["email"] and (r["email"] or "").lower() == v["email"] else "el mismo celular")
            dup.append({"id": r["id"], "nombre": r["nombre_v"], "motivo": motivo,
                        "url": url_for("clientes.detalle", id=r["id"])})
    return v, None, dup


@clientes_bp.route("/crear", methods=["POST"])
@clientes_bp.route("/ajax/crear", methods=["POST"])
@login_required
@tenant_required
def crear():
    tdb = g.tenant_db
    v, err, dup = _leer_cliente(tdb)
    if err:
        return jsonify(ok=False, error=err), 400
    if dup and request.form.get("forzar") != "1":
        return jsonify(ok=False, duplicados=dup, error="Ya hay un cliente parecido"), 409
    if not v["abogado_id"] and session.get("role") in ("abogado", "admin", "super_admin"):
        v["abogado_id"] = session.get("user_id")
    v["creado_por"] = session.get("user_id")
    cur = tdb.execute(f"INSERT INTO clientes ({', '.join(v)}) VALUES ({', '.join('?' * len(v))})", tuple(v.values()))
    tdb.commit()
    cid = cur.lastrowid
    _log("crear_cliente", "cliente", cid, {"nombre": nombre_de(v)})
    destino = url_for("clientes.detalle", id=cid, nuevo=1)
    if request.form.get("volver") == "asunto":
        destino = url_for("asuntos.nuevo", cliente_id=cid)
    return jsonify(ok=True, id=cid, nombre=nombre_de(v), redirect=destino)


@clientes_bp.route("/<int:id>/editar", methods=["POST"])
@login_required
@tenant_required
@role_required("super_admin", "admin", "abogado", "asistente")
def editar(id: int):
    tdb = g.tenant_db
    if not tdb.execute("SELECT 1 FROM clientes WHERE id=? AND activo=1", (id,)).fetchone():
        return jsonify(ok=False, error="Cliente no encontrado"), 404
    v, err, dup = _leer_cliente(tdb, id)
    if err:
        return jsonify(ok=False, error=err), 400
    if dup and request.form.get("forzar") != "1":
        return jsonify(ok=False, duplicados=dup, error="Ya hay un cliente parecido"), 409
    tdb.execute(f"UPDATE clientes SET {', '.join(k + '=?' for k in v)} WHERE id=?", (*v.values(), id))
    tdb.commit()
    _log("actualizar_cliente", "cliente", id, {"campos": "ficha"})
    return jsonify(ok=True)


@clientes_bp.route("/<int:id>/datos")
@login_required
@tenant_required
def datos(id: int):
    r = g.tenant_db.execute("SELECT * FROM clientes WHERE id=? AND activo=1", (id,)).fetchone()
    if not r:
        abort(404)
    d = dict(r)
    d["etiquetas"] = ", ".join(_etiquetas(d.get("etiquetas")))
    return jsonify(d)


_EDITABLE_FIELDS = {"estatus", "abogado_id", "invitar_portal"}


@clientes_bp.route("/<int:id>/editar-campo", methods=["POST"])
@login_required
@tenant_required
def editar_campo(id: int):
    """Cambios rápidos desde la ficha: estatus, responsable y acceso al portal."""
    tdb = g.tenant_db
    if not tdb.execute("SELECT id FROM clientes WHERE id=? AND activo=1", (id,)).fetchone():
        return jsonify(ok=False, error="Cliente no encontrado"), 404
    payload = request.get_json(silent=True) or {}
    campo, valor = payload.get("campo", ""), payload.get("valor")
    if campo not in _EDITABLE_FIELDS:
        return jsonify(ok=False, error="Campo no permitido"), 400
    if campo == "estatus" and valor not in ESTATUS:
        return jsonify(ok=False, error="Estatus no válido"), 400
    if campo == "abogado_id":
        valor = _safe_int(valor, None)
        if valor and valor not in {a["id"] for a in _get_abogados()}:
            return jsonify(ok=False, error="Responsable no válido"), 400
    if campo == "invitar_portal":
        valor = 1 if valor else 0
    tdb.execute(f"UPDATE clientes SET {campo}=? WHERE id=?", (valor, id))
    tdb.commit()
    _log("actualizar_cliente", "cliente", id, {"campo": campo})
    return jsonify(ok=True)


@clientes_bp.route("/crear-rapido", methods=["POST"])
@login_required
@tenant_required
def crear_rapido():
    data = request.get_json(silent=True) or {}
    nombre = (data.get("nombre") or "").strip()[:160]
    if not nombre:
        return jsonify(ok=False, error="Nombre requerido"), 400
    tipo = data.get("tipo") if data.get("tipo") in TIPOS else "persona_fisica"
    cur = g.tenant_db.execute(
        "INSERT INTO clientes (nombre, razon_social, tipo, telefono, email, rfc, activo, creado_por) VALUES (?,?,?,?,?,?,1,?)",
        (nombre, nombre if tipo == "persona_moral" else None, tipo,
         (data.get("telefono") or "").strip() or None, (data.get("email") or "").strip().lower() or None,
         (data.get("rfc") or "").upper().strip() or None, session.get("user_id")))
    g.tenant_db.commit()
    _log("crear_cliente_rapido", "cliente", cur.lastrowid, {"nombre": nombre})
    return jsonify(ok=True, id=cur.lastrowid, nombre=nombre)


@clientes_bp.route("/<int:id>/eliminar", methods=["POST"])
@login_required
@tenant_required
@role_required("super_admin", "admin")
def eliminar(id: int):
    tdb = g.tenant_db
    row = tdb.execute(f"SELECT {NOMBRE_SQL} AS n FROM clientes c WHERE id=? AND activo=1", (id,)).fetchone()
    if not row:
        return jsonify(ok=False, error="No encontrado"), 404
    abiertos = ",".join(f"'{e}'" for e in ABIERTOS)
    en_curso = tdb.execute(f"SELECT COUNT(*) FROM asuntos WHERE cliente_id=? AND estado IN ({abiertos})", (id,)).fetchone()[0]
    if en_curso:
        return jsonify(ok=False, error=f"Tiene {en_curso} asunto{'s' if en_curso > 1 else ''} en curso. Concluye o reasigna sus asuntos antes de eliminarlo."), 400
    tdb.execute("UPDATE clientes SET activo=0 WHERE id=?", (id,))
    tdb.commit()
    _log("eliminar_cliente", "cliente", id, {"nombre": row["n"]})
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return jsonify(ok=True)
    return redirect(url_for("clientes.index"))


# ─────────────────────────────────────────────────────────────────────────────
#  AJAX — abogados y búsqueda rápida (autocomplete para otros módulos)
# ─────────────────────────────────────────────────────────────────────────────


@clientes_bp.route("/ajax/abogados")
@login_required
@tenant_required
def ajax_abogados():
    return jsonify(_get_abogados())


@clientes_bp.route("/ajax/buscar")
@login_required
@tenant_required
def ajax_buscar():
    q = request.args.get("q", "").strip()
    if len(q) < 2:
        return jsonify([])
    like = f"%{q}%"
    rows = g.tenant_db.execute(
        f"""SELECT id, {NOMBRE_SQL} AS label, tipo, email, celular FROM clientes c
            WHERE activo=1 AND ({NOMBRE_SQL} LIKE ? OR rfc LIKE ? OR email LIKE ?)
            ORDER BY label LIMIT 10""", (like, like, like)).fetchall()
    return jsonify([dict(r) for r in rows])


# ─────────────────────────────────────────────────────────────────────────────
#  FICHA DEL CLIENTE
# ─────────────────────────────────────────────────────────────────────────────


def _bitacora(tdb, id: int, usuarios: dict, cobros_ids: list) -> list:
    """Notas, contactos, mensajes del portal y pagos, agrupados por día."""
    items = []
    for n in tdb.execute("SELECT * FROM cliente_notas WHERE cliente_id=?", (id,)).fetchall():
        items.append({"tipo": n["tipo"] if n["tipo"] in BITACORA else "nota", "texto": n["contenido"],
                      "cuando": _local(n["fecha"]), "autor": usuarios.get(n["creado_por"]), "id": n["id"],
                      "borrable": _es_admin() or n["creado_por"] == session.get("user_id")})
    for m in tdb.execute("SELECT * FROM cliente_comunicaciones WHERE cliente_id=?", (id,)).fetchall():
        items.append({"tipo": m["tipo"] if m["tipo"] in BITACORA else "nota", "texto": m["resumen"],
                      "cuando": _local(m["fecha"]), "autor": usuarios.get(m["creado_por"])})
    for p in tdb.execute("SELECT * FROM portal_mensajes WHERE cliente_id=? ORDER BY fecha DESC LIMIT 30", (id,)).fetchall():
        items.append({"tipo": "portal", "texto": p["mensaje"], "cuando": _local(p["fecha"]),
                      "titulo": "Mensaje del cliente en el portal" if p["remitente"] == "cliente" else "Mensaje enviado por el portal"})
    if cobros_ids:
        for p in tdb.execute(
                f"""SELECT p.*, co.concepto, co.moneda FROM cobro_pagos p JOIN cobros co ON co.id=p.cobro_id
                    WHERE p.cobro_id IN ({','.join('?' * len(cobros_ids))})""", cobros_ids).fetchall():
            cuando = _local(p["fecha_alta"])
            f = _fecha(p["fecha"])
            if f and cuando and cuando.date() != f:
                cuando = cuando.replace(year=f.year, month=f.month, day=f.day)
            items.append({"tipo": "pago", "texto": p["concepto"], "cuando": cuando,
                          "titulo": f"Pagó {cobranzas.dinero(p['monto'], p['moneda'])}" + (f" · {p['metodo']}" if p["metodo"] else ""),
                          "autor": usuarios.get(p["creado_por"]), "cobro_id": p["cobro_id"]})
    items = [i for i in items if i["cuando"]]
    items.sort(key=lambda i: i["cuando"], reverse=True)
    hoy = _hoy()
    grupos, actual = [], None
    for i in items[:80]:
        dia = i["cuando"].date()
        etiqueta = ("Hoy" if dia == hoy else "Ayer" if dia == hoy - timedelta(days=1)
                    else f"{dia.day} {MESES[dia.month - 1]}" + (f" {dia.year}" if dia.year != hoy.year else ""))
        if not actual or actual["dia"] != etiqueta:
            actual = {"dia": etiqueta, "items": []}
            grupos.append(actual)
        i["hora"] = i["cuando"].strftime("%H:%M")
        actual["items"].append(i)
    return grupos


@clientes_bp.route("/<int:id>")
@login_required
@tenant_required
def detalle(id: int):
    tdb = g.tenant_db
    row = tdb.execute("SELECT * FROM clientes WHERE id=? AND activo=1", (id,)).fetchone()
    if not row:
        abort(404)
    cobranzas.al_dia(tdb)
    cliente = dict(row)
    cliente["nombre_v"] = nombre_de(cliente)
    cliente["iniciales"], cliente["color"] = _iniciales(cliente["nombre_v"]), _color(id)
    cliente["etiquetas"] = _etiquetas(cliente.get("etiquetas"))
    cliente["estatus_info"] = ESTATUS.get(cliente.get("estatus") or "activo", ESTATUS["activo"])
    usuarios = _usuarios()
    hoy = _hoy()

    # Asuntos (los abogados solo ven los suyos)
    alc, pa = _alc_asuntos("a")
    asuntos = []
    for a in tdb.execute(f"SELECT a.* FROM asuntos a WHERE a.cliente_id=? AND {alc} "
                         f"ORDER BY a.estado IN ('cerrado','completado','archivado'), a.fecha_alta DESC", (id, *pa)).fetchall():
        a = dict(a)
        a["estado_info"] = ESTADOS_ASUNTO.get(a["estado"], {"label": a["estado"], "color": "#60718D", "fondo": "#F1EEE8"})
        a["tipo_info"] = TIPOS_DERECHO.get(a.get("tipo_asunto"), {"label": "General", "color": "#60718D"})
        a["responsable"] = usuarios.get(a.get("abogado_responsable"))
        a["abierto"] = a["estado"] in ABIERTOS
        asuntos.append(a)
    ocultos = 0
    if _solo_lo_suyo():
        ocultos = tdb.execute("SELECT COUNT(*) FROM asuntos WHERE cliente_id=?", (id,)).fetchone()[0] - len(asuntos)

    # Cobranza
    alc_c, pc = cobranzas._alcance()
    cobros = [cobranzas._decorar(r) for r in tdb.execute(
        f"{cobranzas.BASE_SQL} WHERE co.cliente_id=? AND {alc_c} "
        f"ORDER BY co.estado IN ('pagado','cancelado'), date(co.fecha_vencimiento), co.id DESC", (id, *pc)).fetchall()]
    vigentes = [c for c in cobros if c["estado"] != "cancelado"]
    dinero = {
        "facturado": sum(c["total_v"] for c in vigentes),
        "cobrado": sum(c["pagado_v"] for c in vigentes),
        "saldo": sum(c["saldo_v"] for c in vigentes),
        "vencido": sum(c["saldo_v"] for c in vigentes if c["vence"] and (c["vence"]["dias"] or 0) < 0),
    }

    documentos = []
    for d in tdb.execute("SELECT * FROM cliente_documentos WHERE cliente_id=? ORDER BY fecha_subida DESC", (id,)).fetchall():
        d = dict(d)
        d["fecha_txt"] = _fecha_txt(_local(d["fecha_subida"]))
        d["autor"] = usuarios.get(d.get("subido_por"))
        d["ext"] = (d.get("archivo_nombre") or d.get("nombre") or "").rsplit(".", 1)[-1].upper()[:4] if "." in (d.get("archivo_nombre") or "") else "DOC"
        d["borrable"] = _es_admin() or d.get("subido_por") == session.get("user_id")
        documentos.append(d)

    bitacora = _bitacora(tdb, id, usuarios, [c["id"] for c in cobros])
    ultimo = next((i for gr in bitacora for i in gr["items"] if i["tipo"] in BITACORA), None)

    despacho = g.tenant["nombre_despacho"]
    saludo = (cliente["nombre"] or "").split(" ")[0] if cliente["tipo"] == "persona_fisica" else (cliente.get("contacto_principal") or "")
    domicilio = ", ".join(x for x in (
        " ".join(x for x in (cliente.get("direccion"), cliente.get("num_ext") and f"#{cliente['num_ext']}",
                             cliente.get("num_int") and f"int. {cliente['num_int']}") if x),
        cliente.get("colonia") and f"Col. {cliente['colonia']}", cliente.get("municipio") or cliente.get("ciudad"),
        cliente.get("estado"), cliente.get("cp") and f"C.P. {cliente['cp']}") if x)

    alta = _local(cliente.get("fecha_alta"))
    return render_template(
        "clientes/detalle.html",
        cliente=cliente, asuntos=asuntos, ocultos=ocultos, cobros=cobros, dinero=dinero,
        documentos=documentos, bitacora=bitacora, ultimo=ultimo, bitacora_tipos=BITACORA,
        responsable=usuarios.get(cliente.get("abogado_id")), abogados=_get_abogados(),
        whatsapp=_whatsapp(cliente, f"Hola{' ' + saludo if saludo else ''}, le saludamos de {despacho}. "),
        domicilio=domicilio, tipos=TIPOS, estatus=ESTATUS, estados_mx=ESTADOS_MX, origenes=ORIGENES,
        en_curso=sum(1 for a in asuntos if a["abierto"]), alta_txt=_fecha_txt(alta), alta_hace=_hace(alta),
        nacimiento=_fecha_txt(cliente.get("fecha_nacimiento")) if cliente.get("fecha_nacimiento") else None,
        es_admin=_es_admin(), puede_editar=session.get("role") in ("super_admin", "admin", "abogado", "asistente"),
        recien_creado=request.args.get("nuevo") == "1", hoy=hoy.isoformat(),
    )


# ─────────────────────────────────────────────────────────────────────────────
#  BITÁCORA (notas y contactos)
# ─────────────────────────────────────────────────────────────────────────────


@clientes_bp.route("/<int:id>/notas", methods=["POST"])
@login_required
@tenant_required
def agregar_nota(id: int):
    tdb = g.tenant_db
    if not tdb.execute("SELECT 1 FROM clientes WHERE id=? AND activo=1", (id,)).fetchone():
        return jsonify(ok=False, error="Cliente no encontrado"), 404
    data = request.get_json(silent=True) if request.is_json else request.form
    contenido = ((data or {}).get("contenido") or "").strip()[:5000]
    tipo = (data or {}).get("tipo", "nota")
    if not contenido:
        return jsonify(ok=False, error="Escribe algo para guardar en la bitácora"), 400
    if tipo not in TIPOS_NOTA:
        tipo = "nota"
    cur = tdb.execute("INSERT INTO cliente_notas (cliente_id, contenido, tipo, creado_por) VALUES (?,?,?,?)",
                      (id, contenido, tipo, session.get("user_id")))
    tdb.commit()
    return jsonify(ok=True, id=cur.lastrowid)


@clientes_bp.route("/<int:id>/notas/<int:nid>/borrar", methods=["POST"])
@login_required
@tenant_required
def borrar_nota(id: int, nid: int):
    tdb = g.tenant_db
    n = tdb.execute("SELECT creado_por FROM cliente_notas WHERE id=? AND cliente_id=?", (nid, id)).fetchone()
    if not n:
        return jsonify(ok=False, error="No encontrada"), 404
    if not _es_admin() and n["creado_por"] != session.get("user_id"):
        return jsonify(ok=False, error="Solo quien la escribió puede borrarla"), 403
    tdb.execute("DELETE FROM cliente_notas WHERE id=?", (nid,))
    tdb.commit()
    return jsonify(ok=True)


@clientes_bp.route("/<int:id>/comunicaciones", methods=["POST"])
@login_required
@tenant_required
def agregar_comunicacion(id: int):
    data = request.get_json(silent=True) or {}
    resumen = (data.get("resumen") or "").strip()[:5000]
    tipo = data.get("tipo") if data.get("tipo") in TIPOS_COMUNICACION else "llamada"
    if not resumen:
        return jsonify(ok=False, error="Resumen requerido"), 400
    cur = g.tenant_db.execute("INSERT INTO cliente_comunicaciones (cliente_id, tipo, resumen, creado_por) VALUES (?,?,?,?)",
                              (id, tipo, resumen, session.get("user_id")))
    g.tenant_db.commit()
    return jsonify(ok=True, id=cur.lastrowid)


# ─────────────────────────────────────────────────────────────────────────────
#  DOCUMENTOS Y FOTO
# ─────────────────────────────────────────────────────────────────────────────


@clientes_bp.route("/<int:id>/documentos", methods=["POST"])
@login_required
@tenant_required
def subir_documento(id: int):
    tdb = g.tenant_db
    if not tdb.execute("SELECT 1 FROM clientes WHERE id=? AND activo=1", (id,)).fetchone():
        return jsonify(ok=False, error="Cliente no encontrado"), 404
    nombre = (request.form.get("nombre") or "").strip()[:200]
    descripcion = (request.form.get("descripcion") or "").strip()[:500] or None
    f = request.files.get("archivo")
    if not f or not f.filename:
        return jsonify(ok=False, error="Elige el archivo que quieres subir"), 400
    info = boveda.guardar(f, "clientes", id)
    if not info:
        return jsonify(ok=False, error="Ese tipo de archivo no se permite"), 400
    cur = tdb.execute(
        """INSERT INTO cliente_documentos
           (cliente_id, nombre, descripcion, archivo_url, archivo_nombre, tamanio, mime_type, subido_por)
           VALUES (?,?,?,?,?,?,?,?)""",
        (id, nombre or info["nombre"], descripcion, info["url"], info["nombre"], info["tamanio"], info["mime"],
         session.get("user_id")))
    tdb.commit()
    lector.encolar(info["url"], nombre or info["nombre"], "cliente", cur.lastrowid, cliente_id=id)
    return jsonify(ok=True, id=cur.lastrowid)


@clientes_bp.route("/<int:id>/documentos/<int:did>/borrar", methods=["POST"])
@login_required
@tenant_required
def borrar_documento(id: int, did: int):
    tdb = g.tenant_db
    d = tdb.execute("SELECT * FROM cliente_documentos WHERE id=? AND cliente_id=?", (did, id)).fetchone()
    if not d:
        return jsonify(ok=False, error="Documento no encontrado"), 404
    if not _es_admin() and d["subido_por"] != session.get("user_id"):
        return jsonify(ok=False, error="Solo quien lo subió o un administrador puede borrarlo"), 403
    tdb.execute("DELETE FROM cliente_documentos WHERE id=?", (did,))
    tdb.commit()
    if d["archivo_url"]:
        boveda.eliminar(d["archivo_url"])
    _log("borrar_documento_cliente", "cliente", id, {"nombre": d["nombre"]})
    return jsonify(ok=True)


@clientes_bp.route("/<int:id>/foto", methods=["POST"])
@login_required
@tenant_required
def subir_foto(id: int):
    f = request.files.get("foto")
    if not f or not f.filename:
        return jsonify(ok=False, error="Archivo requerido"), 400
    info = boveda.guardar(f, "clientes", "fotos", extensiones={"jpg", "jpeg", "png", "webp", "gif"})
    if not info:
        return jsonify(ok=False, error="La foto debe ser JPG, PNG o WebP"), 400
    anterior = g.tenant_db.execute("SELECT foto_url FROM clientes WHERE id=?", (id,)).fetchone()
    g.tenant_db.execute("UPDATE clientes SET foto_url=? WHERE id=?", (info["url"], id))
    g.tenant_db.commit()
    if anterior and anterior["foto_url"]:
        boveda.eliminar(anterior["foto_url"])
    return jsonify(ok=True, url=info["url"])


# ─────────────────────────────────────────────────────────────────────────────
#  EXCEL: exportar, plantilla e importación
# ─────────────────────────────────────────────────────────────────────────────

_ENCABEZADO = "1D3FD6"


def _encabezados(ws, titulos, anchos):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws.append(titulos)
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor=_ENCABEZADO)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"


@clientes_bp.route("/exportar")
@login_required
@tenant_required
def exportar():
    from openpyxl import Workbook

    tdb = g.tenant_db
    usuarios = _usuarios()
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    _encabezados(ws, ["Nombre o razón social", "Tipo", "Estatus", "RFC", "CURP", "Correo", "Celular", "Teléfono",
                      "Contacto", "Domicilio", "Colonia", "Municipio", "Estado", "C.P.", "Origen", "Responsable",
                      "Etiquetas", "Cliente desde"],
                 [34, 14, 11, 16, 21, 28, 14, 14, 22, 30, 18, 18, 18, 8, 20, 22, 22, 13])
    for r in tdb.execute(f"SELECT c.*, {NOMBRE_SQL} AS nombre_v FROM clientes c WHERE activo=1 ORDER BY nombre_v COLLATE NOCASE"):
        domicilio = " ".join(x for x in (r["direccion"], r["num_ext"] and f"#{r['num_ext']}", r["num_int"] and f"int. {r['num_int']}") if x)
        ws.append([r["nombre_v"], TIPOS.get(r["tipo"], r["tipo"]), ESTATUS.get(r["estatus"] or "activo", {}).get("label", r["estatus"]),
                   r["rfc"], r["curp"], r["email"], r["celular"], r["telefono"], r["contacto_principal"], domicilio,
                   r["colonia"], r["municipio"] or r["ciudad"], r["estado"], r["cp"], r["origen_cliente"],
                   (usuarios.get(r["abogado_id"]) or {}).get("nombre"), ", ".join(_etiquetas(r["etiquetas"])),
                   _fecha(_local(r["fecha_alta"]))])
    for fila in ws.iter_rows(min_row=2, min_col=18, max_col=18):
        fila[0].number_format = "DD/MM/YYYY"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=f"clientes_{_hoy().isoformat()}.xlsx")


_IMPORT_COLS = ["tipo", "nombre", "apellidos", "razon_social", "rfc", "curp", "email", "telefono", "celular",
                "direccion", "colonia", "municipio", "ciudad", "estado", "cp", "origen_cliente", "estatus", "etiquetas"]


@clientes_bp.route("/plantilla-importar")
@login_required
@tenant_required
def plantilla_importar():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    _encabezados(ws, ["Tipo *", "Nombre *", "Apellidos", "Razón social", "RFC", "CURP", "Correo", "Teléfono", "Celular",
                      "Calle y número", "Colonia", "Municipio", "Ciudad", "Estado", "C.P.", "Origen", "Estatus",
                      "Etiquetas (separadas por coma)"],
                 [16, 18, 20, 30, 16, 21, 26, 13, 13, 26, 16, 18, 14, 18, 8, 20, 12, 24])
    ws.append(["persona_fisica o persona_moral", "Nombre de la persona (o de la empresa)", "Solo persona física",
               "Solo empresa", "13 o 12 caracteres", "18 caracteres", "", "10 dígitos", "10 dígitos", "", "", "",
               "", "", "5 dígitos", "", "activo, prospecto o inactivo", "vip, corporativo"])
    for c in ws[2]:
        c.font = Font(italic=True, color="64708A", size=9)
    ws.append(["persona_fisica", "Juan", "Pérez García", "", "PEGJ900101AB1", "", "juan@correo.com", "", "5512345678",
               "Insurgentes Sur 100", "Del Valle", "Benito Juárez", "Ciudad de México", "Ciudad de México", "03100",
               "Referido por cliente", "activo", "vip"])
    dv = DataValidation(type="list", formula1='"persona_fisica,persona_moral"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("A3:A5000")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="plantilla_clientes.xlsx")


def _parse_import_xlsx(file_obj) -> tuple[list[dict], list[dict]]:
    from openpyxl import load_workbook
    ws = load_workbook(file_obj, data_only=True, read_only=True).active
    ok, errores = [], []
    for fila, row in enumerate(ws.iter_rows(min_row=3, max_col=len(_IMPORT_COLS), values_only=True), start=3):
        if not any(row):
            continue
        r = {k: (str(row[i]).strip() if i < len(row) and row[i] is not None else "") for i, k in enumerate(_IMPORT_COLS)}
        if r["nombre"].startswith("Nombre de la persona"):
            continue
        tipo = r["tipo"] if r["tipo"] in TIPOS else ("persona_moral" if r["razon_social"] and not r["apellidos"] else "persona_fisica")
        rfc = re.sub(r"[\s-]", "", r["rfc"].upper())
        problemas = []
        if not (r["nombre"] or r["razon_social"]):
            problemas.append("falta el nombre")
        if rfc and not RE_RFC.match(rfc):
            problemas.append("RFC con formato inválido")
        if r["email"] and not RE_EMAIL.match(r["email"]):
            problemas.append("correo inválido")
        registro = {
            "fila": fila, "tipo": tipo,
            "nombre": (r["razon_social"] or r["nombre"]) if tipo == "persona_moral" else r["nombre"],
            "apellidos": None if tipo == "persona_moral" else (r["apellidos"] or None),
            "razon_social": (r["razon_social"] or r["nombre"]) if tipo == "persona_moral" else None,
            "rfc": rfc or None, "curp": r["curp"].upper() or None, "email": r["email"].lower() or None,
            "telefono": r["telefono"] or None, "celular": r["celular"] or None, "direccion": r["direccion"] or None,
            "colonia": r["colonia"] or None, "municipio": r["municipio"] or None, "ciudad": r["ciudad"] or None,
            "estado": r["estado"] or None, "cp": r["cp"][:5] or None, "origen_cliente": r["origen_cliente"] or None,
            "estatus": r["estatus"].lower() if r["estatus"].lower() in ESTATUS else "activo",
            "etiquetas": json.dumps([t.strip() for t in r["etiquetas"].split(",") if t.strip()][:12], ensure_ascii=False),
        }
        if problemas:
            registro["errores"] = problemas
            errores.append(registro)
        else:
            ok.append(registro)
    return ok, errores


@clientes_bp.route("/importar/preview", methods=["POST"])
@login_required
@tenant_required
@role_required("super_admin", "admin", "abogado", "asistente")
def importar_preview():
    f = request.files.get("archivo")
    if not f or not f.filename:
        return jsonify(ok=False, error="Elige el archivo de Excel"), 400
    try:
        rows_ok, rows_err = _parse_import_xlsx(f)
    except Exception:
        return jsonify(ok=False, error="No se pudo leer el archivo. Usa la plantilla de Excel (.xlsx)."), 400
    existentes = {r[0] for r in g.tenant_db.execute("SELECT rfc FROM clientes WHERE activo=1 AND rfc IS NOT NULL")}
    for r in rows_ok:
        r["repetido"] = bool(r["rfc"] and r["rfc"] in existentes)
    return jsonify(ok=True, rows_ok=rows_ok, rows_err=rows_err, total_ok=len(rows_ok), total_err=len(rows_err))


@clientes_bp.route("/importar/confirmar", methods=["POST"])
@login_required
@tenant_required
@role_required("super_admin", "admin", "abogado", "asistente")
def importar_confirmar():
    rows = (request.get_json(silent=True) or {}).get("rows", [])[:5000]
    if not rows:
        return jsonify(ok=False, error="Sin filas para importar"), 400
    tdb = g.tenant_db
    cols = [c for c in _IMPORT_COLS]
    insertados = 0
    for r in rows:
        if not r.get("nombre") or r.get("repetido"):
            continue
        r["tipo"] = r.get("tipo") if r.get("tipo") in TIPOS else "persona_fisica"
        r["estatus"] = r.get("estatus") if r.get("estatus") in ESTATUS else "activo"
        try:
            tdb.execute(f"INSERT INTO clientes ({', '.join(cols)}, creado_por) VALUES ({', '.join('?' * len(cols))}, ?)",
                        (*[(str(r.get(c))[:300] if r.get(c) is not None else None) for c in cols], session.get("user_id")))
            insertados += 1
        except Exception:
            pass
    tdb.commit()
    _log("importar_clientes", "cliente", detalles={"insertados": insertados})
    return jsonify(ok=True, insertados=insertados)


# ─────────────────────────────────────────────────────────────────────────────
#  FICHA EN PDF
# ─────────────────────────────────────────────────────────────────────────────


@clientes_bp.route("/<int:id>/pdf")
@login_required
@tenant_required
def exportar_pdf(id: int):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    tdb = g.tenant_db
    row = tdb.execute("SELECT * FROM clientes WHERE id=? AND activo=1", (id,)).fetchone()
    if not row:
        abort(404)
    c = dict(row)
    nombre = nombre_de(c)
    usuarios = _usuarios()
    alc, pa = _alc_asuntos("a")
    asuntos = tdb.execute(f"SELECT a.* FROM asuntos a WHERE a.cliente_id=? AND {alc} ORDER BY a.fecha_alta DESC LIMIT 20",
                          (id, *pa)).fetchall()
    alc_c, pc = cobranzas._alcance()
    cobros = [cobranzas._decorar(r) for r in tdb.execute(
        f"{cobranzas.BASE_SQL} WHERE co.cliente_id=? AND co.estado!='cancelado' AND {alc_c}", (id, *pc)).fetchall()]

    AZUL, TINTA, GRIS, LINEA = (colors.HexColor(x) for x in ("#142E59", "#142E59", "#60718D", "#E5E2DB"))
    e = lambda x: html.escape(str(x or "—"))  # noqa: E731
    st = lambda n, **k: ParagraphStyle(n, fontName=k.pop("font", "Helvetica"), **k)  # noqa: E731
    s_h = st("h", font="Helvetica-Bold", fontSize=10.5, textColor=TINTA, spaceBefore=14, spaceAfter=6)
    s_c = st("c", fontSize=8.5, textColor=TINTA, leading=11)
    despacho = g.tenant["nombre_despacho"]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.6 * cm,
                            bottomMargin=1.6 * cm, title=f"Ficha de cliente — {nombre}")
    story = [Paragraph(e(despacho), st("d", font="Helvetica-Bold", fontSize=10, textColor=AZUL)), Spacer(1, 6),
             Paragraph(e(nombre), st("t", font="Helvetica-Bold", fontSize=18, textColor=TINTA, leading=22)),
             Paragraph(f"{TIPOS.get(c['tipo'], '')} · {ESTATUS.get(c.get('estatus') or 'activo', {}).get('label', '')}",
                       st("s", fontSize=9, textColor=GRIS, leading=13))]

    def tabla_kv(filas):
        filas = [[k, Paragraph(e(v), s_c)] for k, v in filas if v]
        if not filas:
            return
        t = Table(filas, colWidths=[4.4 * cm, 13.2 * cm])
        t.setStyle(TableStyle([("FONT", (0, 0), (0, -1), "Helvetica", 8.5), ("TEXTCOLOR", (0, 0), (0, -1), GRIS),
                               ("LINEBELOW", (0, 0), (-1, -1), .5, LINEA), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        story.append(t)

    story.append(Paragraph("Datos", s_h))
    domicilio = ", ".join(x for x in (c.get("direccion"), c.get("num_ext") and f"#{c['num_ext']}", c.get("colonia"),
                                       c.get("municipio") or c.get("ciudad"), c.get("estado"), c.get("cp") and f"C.P. {c['cp']}") if x)
    tabla_kv([("RFC", c.get("rfc")), ("CURP", c.get("curp")), ("Correo", c.get("email")), ("Celular", c.get("celular")),
              ("Teléfono", c.get("telefono")), ("Contacto", c.get("contacto_principal")),
              ("Representante legal", c.get("representante_legal")), ("Domicilio", domicilio),
              ("Responsable", (usuarios.get(c.get("abogado_id")) or {}).get("nombre")), ("Origen", c.get("origen_cliente")),
              ("Cliente desde", _fecha_txt(_local(c.get("fecha_alta"))))])

    if asuntos:
        story.append(Paragraph("Asuntos", s_h))
        filas = [["Asunto", "Expediente", "Materia", "Estado"]]
        for a in asuntos:
            filas.append([Paragraph(e(a["titulo"]), s_c), a["numero_expediente"] or "—",
                          TIPOS_DERECHO.get(a["tipo_asunto"], {}).get("label", "—"),
                          ESTADOS_ASUNTO.get(a["estado"], {}).get("label", a["estado"])])
        t = Table(filas, colWidths=[8.6 * cm, 3 * cm, 3 * cm, 3 * cm], repeatRows=1)
        t.setStyle(TableStyle([("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.5), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("BACKGROUND", (0, 0), (-1, 0), TINTA), ("FONT", (1, 1), (-1, -1), "Helvetica", 8),
                               ("LINEBELOW", (0, 1), (-1, -1), .5, LINEA), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                               ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        story.append(t)

    if cobros:
        story.append(Paragraph("Cobranza", s_h))
        tabla_kv([("Facturado", cobranzas.dinero(sum(x["total_v"] for x in cobros))),
                  ("Pagado", cobranzas.dinero(sum(x["pagado_v"] for x in cobros))),
                  ("Saldo pendiente", cobranzas.dinero(sum(x["saldo_v"] for x in cobros)))])

    story += [Spacer(1, 18), Paragraph(f"{e(despacho)} · Documento confidencial · Generado el {_fecha_txt(_hoy())} con Habeas Corpus",
                                       st("p", fontSize=7.5, textColor=GRIS))]
    doc.build(story)
    buf.seek(0)
    seguro = re.sub(r"[^A-Za-z0-9]+", "_", nombre).strip("_")[:40] or "cliente"
    return send_file(buf, mimetype="application/pdf", as_attachment=True, download_name=f"ficha_{seguro}.pdf")
