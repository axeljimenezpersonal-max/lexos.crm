"""
Documentos del editor a Word y PDF, conservando lo que se ve en pantalla: títulos, negritas, cursivas,
subrayado, listas, sangrías y alineación (incluido justificado). Letra Arial 12, como se presenta en México.
"""
from __future__ import annotations

import io
import re
from html.parser import HTMLParser

_BLOQUES = {"p", "h1", "h2", "h3", "h4", "li", "blockquote", "div"}
_ALINEA = re.compile(r"ql-align-(center|right|justify)")
_SANGRIA = re.compile(r"ql-indent-(\d)")


class _Lector(HTMLParser):
    """Convierte el HTML del editor (Quill) en bloques con trozos de texto y su formato."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.bloques: list[dict] = []
        self.actual: dict | None = None
        self.listas: list[str] = []
        self.fmt = {"b": 0, "i": 0, "u": 0}

    def _nuevo(self, tipo: str, attrs: dict):
        clase = attrs.get("class") or ""
        alinea = _ALINEA.search(clase)
        sangria = _SANGRIA.search(clase)
        self.actual = {"tipo": tipo, "alinea": alinea.group(1) if alinea else "left",
                       "sangria": int(sangria.group(1)) if sangria else 0, "trozos": []}
        if tipo == "li":
            self.actual["lista"] = self.listas[-1] if self.listas else "ul"
        self.bloques.append(self.actual)

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in ("ol", "ul"):
            self.listas.append(tag)
        elif tag in _BLOQUES:
            self._nuevo(tag if tag != "div" else "p", attrs)
        elif tag in ("strong", "b"):
            self.fmt["b"] += 1
        elif tag in ("em", "i"):
            self.fmt["i"] += 1
        elif tag == "u":
            self.fmt["u"] += 1
        elif tag == "br" and self.actual is not None:
            self.actual["trozos"].append(("\n", False, False, False))

    def handle_endtag(self, tag):
        if tag in ("ol", "ul") and self.listas:
            self.listas.pop()
        elif tag in _BLOQUES:
            self.actual = None
        elif tag in ("strong", "b"):
            self.fmt["b"] = max(0, self.fmt["b"] - 1)
        elif tag in ("em", "i"):
            self.fmt["i"] = max(0, self.fmt["i"] - 1)
        elif tag == "u":
            self.fmt["u"] = max(0, self.fmt["u"] - 1)

    def handle_data(self, data):
        if not data:
            return
        if self.actual is None:
            if not data.strip():
                return
            self._nuevo("p", {})
        self.actual["trozos"].append((data, bool(self.fmt["b"]), bool(self.fmt["i"]), bool(self.fmt["u"])))


def bloques(html: str) -> list[dict]:
    lector = _Lector()
    lector.feed(html or "")
    lector.close()
    for b in lector.bloques:   # un párrafo que solo tiene un salto de línea es un renglón vacío
        if all(t[0] == "\n" for t in b["trozos"]):
            b["trozos"] = []
    return lector.bloques


def _numeros(bs: list[dict]) -> list[int]:
    """Número de cada elemento de lista numerada (se reinicia al cambiar de lista)."""
    salida, n, previo = [], 0, None
    for b in bs:
        if b["tipo"] == "li" and b.get("lista") == "ol":
            n = n + 1 if previo == "ol" else 1
        salida.append(n)
        previo = b.get("lista") if b["tipo"] == "li" else None
    return salida


# ── Word ──────────────────────────────────────────────────────────────────────

def a_docx(html: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    alinea = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
              "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin = s.bottom_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5

    bs = bloques(html)
    for b, num in zip(bs, _numeros(bs)):
        p = doc.add_paragraph()
        p.alignment = alinea[b["alinea"]]
        fmt = p.paragraph_format
        fmt.left_indent = Cm(1.25 * b["sangria"] + (0.75 if b["tipo"] in ("li", "blockquote") else 0))
        tam, negrita = {"h1": (16, True), "h2": (14, True), "h3": (12.5, True), "h4": (12, True)}.get(b["tipo"], (12, False))
        if b["tipo"] == "li":
            p.add_run(f"{num}. " if b.get("lista") == "ol" else "•  ")
        for texto, bo, it, su in b["trozos"]:
            if texto == "\n":
                p.add_run().add_break()
                continue
            r = p.add_run(texto)
            r.bold, r.italic, r.underline = (bo or negrita), (it or b["tipo"] == "blockquote"), su
            r.font.size = Pt(tam)
            if b["tipo"] in ("h1", "h2", "h3", "h4"):
                r.font.color.rgb = RGBColor(0x14, 0x2E, 0x59)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── PDF ───────────────────────────────────────────────────────────────────────

def _xml(texto: str) -> str:
    return texto.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def a_pdf(html: str, titulo: str = "Documento") -> bytes:
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    alinea = {"left": TA_LEFT, "center": TA_CENTER, "right": TA_RIGHT, "justify": TA_JUSTIFY}
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=12, leading=18, spaceAfter=6)
    tamanos = {"h1": 16, "h2": 14, "h3": 12.5, "h4": 12}
    historia = []
    bs = bloques(html)
    for b, num in zip(bs, _numeros(bs)):
        if not b["trozos"]:
            historia.append(Spacer(1, 10))
            continue
        partes = []
        for texto, bo, it, su in b["trozos"]:
            if texto == "\n":
                partes.append("<br/>")
                continue
            t = _xml(texto)
            if su:
                t = f"<u>{t}</u>"
            if it or b["tipo"] == "blockquote":
                t = f"<i>{t}</i>"
            if bo or b["tipo"] in tamanos:
                t = f"<b>{t}</b>"
            partes.append(t)
        sangria = 1.25 * cm * b["sangria"] + (0.75 * cm if b["tipo"] in ("li", "blockquote") else 0)
        estilo = ParagraphStyle(f"b{len(historia)}", parent=base, alignment=alinea[b["alinea"]], leftIndent=sangria,
                                fontSize=tamanos.get(b["tipo"], 12), leading=tamanos.get(b["tipo"], 12) * 1.5,
                                textColor="#142E59" if b["tipo"] in tamanos else "#1A1A1A",
                                spaceBefore=8 if b["tipo"] in tamanos else 0)
        texto = "".join(partes)
        if b["tipo"] == "li":
            texto = (f"{num}. " if b.get("lista") == "ol" else "•&nbsp;&nbsp;") + texto
        historia.append(Paragraph(texto, estilo))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=2.5 * cm, rightMargin=2.5 * cm,
                      topMargin=2.5 * cm, bottomMargin=2.5 * cm, title=titulo).build(historia or [Spacer(1, 1)])
    return buf.getvalue()
