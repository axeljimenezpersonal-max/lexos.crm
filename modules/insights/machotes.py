"""
Machotes: plantillas Word (.docx) con campos {{así}} que se llenan con los datos del asunto.

Todo el llenado ocurre en el servidor del despacho: el machote y los datos nunca se mandan
a ningún modelo ni servicio externo.
"""
from __future__ import annotations

import html
import io
import re
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

PATRON = re.compile(r"\{\{\s*([A-Za-z_ÁÉÍÓÚÑáéíóúñ0-9]+)\s*\}\}")
EN_BLANCO = "________________"

# Campos que el sistema sabe llenar solo. El orden es el que se muestra en pantalla.
CAMPOS = [
    ("nombre_cliente", "Nombre del cliente"),
    ("representante_legal", "Representante legal del cliente"),
    ("rfc_cliente", "RFC del cliente"),
    ("curp_cliente", "CURP del cliente"),
    ("domicilio_cliente", "Domicilio del cliente"),
    ("email_cliente", "Correo del cliente"),
    ("telefono_cliente", "Teléfono del cliente"),
    ("titulo_asunto", "Título del asunto"),
    ("numero_expediente", "Número de expediente"),
    ("numero_juicio", "Número de juicio"),
    ("juzgado", "Juzgado"),
    ("materia", "Materia"),
    ("contraparte", "Contraparte"),
    ("abogado_contraparte", "Abogado de la contraparte"),
    ("monto", "Monto del asunto"),
    ("monto_letra", "Monto del asunto con letra"),
    ("honorarios", "Honorarios pactados"),
    ("honorarios_letra", "Honorarios con letra"),
    ("fecha_inicio_asunto", "Fecha de inicio del asunto"),
    ("nombre_abogado", "Abogado que firma"),
    ("despacho", "Nombre del despacho"),
    ("domicilio_despacho", "Domicilio del despacho"),
    ("telefono_despacho", "Teléfono del despacho"),
    ("rfc_despacho", "RFC del despacho"),
    ("fecha_hoy", "Fecha de hoy (DD/MM/AAAA)"),
    ("fecha_hoy_letra", "Fecha de hoy con letra"),
]
ETIQUETAS = dict(CAMPOS)
# Nombres que ya usaba el editor de documentos
ALIAS = {"direccion_cliente": "domicilio_cliente"}

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]


# ── Números y fechas en letra ─────────────────────────────────────────────────

_UNIDADES = ["", "UN", "DOS", "TRES", "CUATRO", "CINCO", "SEIS", "SIETE", "OCHO", "NUEVE", "DIEZ",
             "ONCE", "DOCE", "TRECE", "CATORCE", "QUINCE", "DIECISÉIS", "DIECISIETE", "DIECIOCHO",
             "DIECINUEVE", "VEINTE", "VEINTIÚN", "VEINTIDÓS", "VEINTITRÉS", "VEINTICUATRO", "VEINTICINCO",
             "VEINTISÉIS", "VEINTISIETE", "VEINTIOCHO", "VEINTINUEVE"]
_DECENAS = ["", "", "", "TREINTA", "CUARENTA", "CINCUENTA", "SESENTA", "SETENTA", "OCHENTA", "NOVENTA"]
_CENTENAS = ["", "CIENTO", "DOSCIENTOS", "TRESCIENTOS", "CUATROCIENTOS", "QUINIENTOS", "SEISCIENTOS",
             "SETECIENTOS", "OCHOCIENTOS", "NOVECIENTOS"]


def _cientos(n: int) -> str:
    if n == 100:
        return "CIEN"
    c, resto = divmod(n, 100)
    if resto < 30:
        texto = _UNIDADES[resto]
    else:
        d, u = divmod(resto, 10)
        texto = _DECENAS[d] + (f" Y {_UNIDADES[u]}" if u else "")
    return " ".join(p for p in (_CENTENAS[c], texto) if p)


def _entero_letra(n: int) -> str:
    if n == 0:
        return "CERO"
    partes = []
    millones, resto = divmod(n, 1_000_000)
    miles, unidades = divmod(resto, 1000)
    if millones:
        partes.append("UN MILLÓN" if millones == 1 else f"{_entero_letra(millones)} MILLONES")
    if miles:
        partes.append("MIL" if miles == 1 else f"{_cientos(miles)} MIL")
    if unidades:
        partes.append(_cientos(unidades))
    return " ".join(partes)


def pesos_letra(monto) -> str:
    """50000 → "CINCUENTA MIL PESOS 00/100 M.N.", como se escribe en contratos y demandas."""
    try:
        valor = round(float(monto), 2)
    except (TypeError, ValueError):
        return ""
    enteros = int(valor)
    centavos = int(round((valor - enteros) * 100))
    texto = _entero_letra(enteros)
    de = " DE" if enteros >= 1_000_000 and enteros % 1_000_000 == 0 else ""
    return f"{texto}{de} {'PESO' if enteros == 1 else 'PESOS'} {centavos:02d}/100 M.N."


def fecha_letra(d: date) -> str:
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def _moneda(v) -> str:
    try:
        return f"${float(v):,.2f}"
    except (TypeError, ValueError):
        return ""


# ── Datos ─────────────────────────────────────────────────────────────────────

def _txt(v) -> str:
    return str(v).strip() if v not in (None, "") else ""


def datos(tdb, tenant, user, asunto_id=None, cliente_id=None) -> dict[str, str]:
    """Valor de cada campo conocido para un asunto (y su cliente). Lo que no exista queda vacío."""
    hoy = date.today()
    valores = {clave: "" for clave, _ in CAMPOS}
    valores.update({
        "fecha_hoy": hoy.strftime("%d/%m/%Y"),
        "fecha_hoy_letra": fecha_letra(hoy),
        "despacho": _txt(tenant["nombre_despacho"]) if tenant else "",
        "domicilio_despacho": _txt(tenant["direccion"]) if tenant else "",
        "telefono_despacho": _txt(tenant["telefono"]) if tenant else "",
        "rfc_despacho": _txt(tenant["rfc"]) if tenant else "",
        "nombre_abogado": f"{user['nombre']} {user['apellidos']}".strip() if user else "",
    })

    asunto = tdb.execute("SELECT * FROM asuntos WHERE id=?", (asunto_id,)).fetchone() if asunto_id else None
    if asunto:
        cliente_id = cliente_id or asunto["cliente_id"]
        valores.update({
            "titulo_asunto": _txt(asunto["titulo"]),
            "numero_expediente": _txt(asunto["numero_expediente"]),
            "numero_juicio": _txt(asunto["numero_juicio"]),
            "juzgado": _txt(asunto["juzgado"]),
            "materia": _txt(asunto["tipo_asunto"]).capitalize(),
            "contraparte": _txt(asunto["contrario_nombre"]),
            "abogado_contraparte": _txt(asunto["contrario_abogado"]),
            "fecha_inicio_asunto": _txt(asunto["fecha_inicio"])[:10],
        })
        if asunto["monto_estimado"]:
            valores["monto"] = _moneda(asunto["monto_estimado"])
            valores["monto_letra"] = pesos_letra(asunto["monto_estimado"])
        if asunto["honorarios_pactados"]:
            valores["honorarios"] = _moneda(asunto["honorarios_pactados"])
            valores["honorarios_letra"] = pesos_letra(asunto["honorarios_pactados"])
        if asunto["abogado_responsable"] and tenant is not None:
            from flask import g
            abogado = g.global_db.execute(
                "SELECT nombre, apellidos FROM tenant_users WHERE id=? AND tenant_id=?",
                (asunto["abogado_responsable"], tenant["id"]),
            ).fetchone()
            if abogado:
                valores["nombre_abogado"] = f"{abogado['nombre']} {abogado['apellidos']}".strip()

    cliente = tdb.execute("SELECT * FROM clientes WHERE id=?", (cliente_id,)).fetchone() if cliente_id else None
    if cliente:
        domicilio = ", ".join(p for p in (
            " ".join(p for p in (_txt(cliente["direccion"]), _txt(cliente["num_ext"])) if p),
            f"int. {cliente['num_int']}" if cliente["num_int"] else "",
            f"col. {cliente['colonia']}" if cliente["colonia"] else "",
            _txt(cliente["municipio"]) or _txt(cliente["ciudad"]),
            _txt(cliente["estado"]),
            f"C.P. {cliente['cp']}" if cliente["cp"] else "",
        ) if p)
        valores.update({
            "nombre_cliente": _txt(cliente["razon_social"]) or
                              f"{_txt(cliente['nombre'])} {_txt(cliente['apellidos'])}".strip(),
            "representante_legal": _txt(cliente["representante_legal"]),
            "rfc_cliente": _txt(cliente["rfc"]),
            "curp_cliente": _txt(cliente["curp"]),
            "domicilio_cliente": domicilio,
            "email_cliente": _txt(cliente["email"]),
            "telefono_cliente": _txt(cliente["telefono"]) or _txt(cliente["celular"]),
        })
    return valores


def resolver_html(contenido: str, valores: dict[str, str]) -> str:
    """Reemplaza {{campos}} en HTML del editor, escapando los valores."""
    def cambio(m):
        clave = ALIAS.get(m.group(1).lower(), m.group(1).lower())
        if clave not in valores:
            return m.group(0)
        return html.escape(valores[clave]) if valores[clave] else EN_BLANCO
    return PATRON.sub(cambio, contenido)


# ── Word ──────────────────────────────────────────────────────────────────────

def _parrafos(doc):
    """Todos los párrafos del documento: cuerpo, tablas (anidadas), encabezados y pies."""
    def de_contenedor(cont):
        yield from cont.paragraphs
        for tabla in getattr(cont, "tables", []):
            for fila in tabla.rows:
                for celda in fila.cells:
                    yield from de_contenedor(celda)
    yield from de_contenedor(doc)
    for seccion in doc.sections:
        for parte in (seccion.header, seccion.footer, seccion.first_page_header, seccion.first_page_footer):
            if parte is not None and not parte.is_linked_to_previous:
                yield from de_contenedor(parte)


def campos(contenido: bytes) -> list[dict]:
    """Campos {{...}} que trae el machote, en orden de aparición."""
    doc = Document(io.BytesIO(contenido))
    vistos: list[str] = []
    for p in _parrafos(doc):
        for m in PATRON.finditer(p.text):
            clave = ALIAS.get(m.group(1).lower(), m.group(1).lower())
            if clave not in vistos:
                vistos.append(clave)
    return [{"clave": c, "etiqueta": ETIQUETAS.get(c) or c.replace("_", " ").capitalize(),
             "conocido": c in ETIQUETAS} for c in vistos]


def _reemplazar(parrafo, valores: dict[str, str]) -> None:
    """Reemplaza campos aunque Word los haya partido en varios fragmentos con formato distinto.
    El valor toma el formato del fragmento donde empieza el campo."""
    runs = parrafo.runs
    texto = "".join(r.text for r in runs)
    for m in reversed(list(PATRON.finditer(texto))):
        clave = ALIAS.get(m.group(1).lower(), m.group(1).lower())
        if clave not in valores:
            continue
        valor = valores[clave] or EN_BLANCO
        limites, pos = [], 0
        for r in runs:
            limites.append((pos, pos + len(r.text)))
            pos += len(r.text)
        i = next(k for k, (a, b) in enumerate(limites) if a <= m.start() < b)
        j = next(k for k, (a, b) in enumerate(limites) if a < m.end() <= b)
        ini_i, ini_j = limites[i][0], limites[j][0]
        if i == j:
            t = runs[i].text
            runs[i].text = t[:m.start() - ini_i] + valor + t[m.end() - ini_i:]
        else:
            runs[i].text = runs[i].text[:m.start() - ini_i] + valor
            for k in range(i + 1, j):
                runs[k].text = ""
            runs[j].text = runs[j].text[m.end() - ini_j:]


def llenar(contenido: bytes, valores: dict[str, str]) -> bytes:
    doc = Document(io.BytesIO(contenido))
    for p in _parrafos(doc):
        if "{{" in p.text:
            _reemplazar(p, valores)
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


_ALINEACION = {WD_ALIGN_PARAGRAPH.CENTER: "center", WD_ALIGN_PARAGRAPH.RIGHT: "right",
               WD_ALIGN_PARAGRAPH.JUSTIFY: "justify"}


def _parrafo_html(p) -> str:
    # Se juntan los fragmentos seguidos con el mismo formato: Word suele partir un {{campo}} en varios.
    grupos: list[list] = []
    for r in p.runs:
        if not r.text:
            continue
        formato = (bool(r.bold), bool(r.italic), bool(r.underline))
        if grupos and grupos[-1][0] == formato:
            grupos[-1][1] += r.text
        else:
            grupos.append([formato, r.text])
    trozos = []
    for (negrita, cursiva, subrayado), texto in grupos:
        t = html.escape(texto)
        if negrita:
            t = f"<strong>{t}</strong>"
        if cursiva:
            t = f"<em>{t}</em>"
        if subrayado:
            t = f"<u>{t}</u>"
        trozos.append(t)
    cuerpo = "".join(trozos) or "<br>"
    estilo = (p.style.name or "").lower() if p.style is not None else ""
    alinear = _ALINEACION.get(p.alignment)
    attr = f' style="text-align:{alinear}"' if alinear else ""
    if estilo.startswith("heading 1") or estilo == "title":
        return f"<h2{attr}>{cuerpo}</h2>"
    if estilo.startswith("heading"):
        return f"<h3{attr}>{cuerpo}</h3>"
    return f"<p{attr}>{cuerpo}</p>"


def a_html(contenido: bytes) -> str:
    """Versión editable del machote para el editor de documentos (texto, negritas, tablas)."""
    doc = Document(io.BytesIO(contenido))
    partes = []
    cuerpo = doc.element.body
    for hijo in cuerpo.iterchildren():
        etiqueta = hijo.tag.rsplit("}", 1)[-1]
        if etiqueta == "p":
            from docx.text.paragraph import Paragraph
            partes.append(_parrafo_html(Paragraph(hijo, doc)))
        elif etiqueta == "tbl":
            from docx.table import Table
            tabla = Table(hijo, doc)
            filas = "".join(
                "<tr>" + "".join(f"<td>{''.join(_parrafo_html(p) for p in c.paragraphs)}</td>" for c in f.cells) + "</tr>"
                for f in tabla.rows)
            partes.append(f'<table border="1" style="border-collapse:collapse;width:100%">{filas}</table>')
    return "\n".join(partes)
