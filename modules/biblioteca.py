"""
LexOS — modules/biblioteca.py
Biblioteca: Estudio (archivos), Formatos (plantillas), Acuerdos/Audiencias.
Editor de documentos WYSIWYG con variables dinámicas, versiones, export PDF/DOCX.
"""
from __future__ import annotations

import io
import json
import math
import os
import re
import secrets
import unicodedata
from datetime import date, datetime, timedelta

from flask import (
    Blueprint, Response, abort, current_app, g, jsonify, make_response, redirect, render_template,
    request, send_file, session, url_for,
)
from markupsafe import Markup, escape

from modules import boveda
from modules import lector
from modules import tipos_documento as tipos
from modules.insights import machotes
from modules.auth import login_required, role_required, tenant_required

biblioteca_bp = Blueprint("biblioteca", __name__)

# ─────────────────────────────────────────────────────────────────────────────
#  Constantes
# ─────────────────────────────────────────────────────────────────────────────

CATEGORIAS_PLANTILLA = [
    "Contratos", "Demandas", "Actas", "Poderes", "Convenios",
    "Amparos", "Recursos", "Escritos", "Oficios", "Otro",
]

TIPOS_DERECHO = [
    "Civil", "Penal", "Mercantil", "Laboral", "Familiar",
    "Administrativo", "Fiscal", "Amparo", "Corporativo", "Otro",
]

PLAN_LIMITS_MB = {
    "starter": 500,
    "basico": 500,
    "pro": 2048,
    "profesional": 2048,
    "enterprise": 10240,
}

ALLOWED_EXT = {
    "pdf", "doc", "docx", "xls", "xlsx", "ppt", "pptx",
    "png", "jpg", "jpeg", "gif", "webp", "svg",
    "txt", "csv", "xml", "zip", "rar",
}

VARIABLES_DINAMICAS = [("{{%s}}" % clave, etiqueta) for clave, etiqueta in machotes.CAMPOS]

PER_PAGE = 20

# ─────────────────────────────────────────────────────────────────────────────
#  Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _log(accion, entidad_id=None, detalles=None):
    try:
        g.tenant_db.execute(
            "INSERT INTO actividad_log (user_id, accion, entidad, entidad_id, detalles, ip) VALUES (?,?,?,?,?,?)",
            (session.get("user_id"), accion, "biblioteca", entidad_id,
             json.dumps(detalles or {}), request.remote_addr),
        )
        g.tenant_db.commit()
    except Exception:
        pass


def _safe_int(v, default=None):
    try:
        return int(v) if v not in (None, "", "None") else default
    except (TypeError, ValueError):
        return default


def _save_file(f, subdir):
    """Guarda el archivo cifrado en la bóveda del despacho."""
    info = boveda.guardar(f, "biblioteca", subdir, extensiones=ALLOWED_EXT)
    if not info:
        return None
    return {"url": info["url"], "nombre_original": info["nombre"], "tamanio": info["tamanio"],
            "mime_type": info["mime"], "extension": info["ext"]}


def _espacio_usado(tdb) -> int:
    """Total bytes used by all archivos in this tenant's estudio."""
    r = tdb.execute("SELECT COALESCE(SUM(tamanio), 0) FROM biblioteca_archivos").fetchone()
    return r[0]


def _plan_limit_bytes():
    plan = g.tenant["plan"] if g.tenant else "starter"
    mb = PLAN_LIMITS_MB.get(plan, 500)
    return mb * 1024 * 1024


def _get_abogados():
    try:
        return [dict(r) for r in g.global_db.execute(
            "SELECT id, nombre, apellidos FROM tenant_users "
            "WHERE tenant_id=? AND rol IN ('abogado','admin','super_admin') AND activo=1 ORDER BY nombre",
            (g.tenant["id"],),
        ).fetchall()]
    except Exception:
        return []


def _resolve_variables(html: str, asunto_id=None, cliente_id=None) -> str:
    """Reemplaza {{variables}} con los datos del asunto, del cliente y del despacho."""
    valores = machotes.datos(g.tenant_db, g.tenant, g.user, asunto_id, cliente_id)
    return machotes.resolver_html(html, valores)


# ═════════════════════════════════════════════════════════════════════════════
#  INDEX — página principal con 3 secciones (tabs)
# ═════════════════════════════════════════════════════════════════════════════


_MESES_CORTOS = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]

# Estado de lectura de cada documento (OCR e índice), con texto y no solo color
ESTADOS_LECTURA = {
    "listo":       {"label": "Listo",             "grupo": "listo"},
    "pendiente":   {"label": "Procesando",        "grupo": "procesando"},
    "procesando":  {"label": "Procesando",        "grupo": "procesando"},
    "requiere_ocr": {"label": "Requiere revisión", "grupo": "revision"},
    "sin_texto":   {"label": "Requiere revisión", "grupo": "revision"},
    "no_soportado": {"label": "Requiere revisión", "grupo": "revision"},
    "error":       {"label": "Error",             "grupo": "error"},
}


def _documentos_de_expedientes(tdb):
    """Documentos y resoluciones de los expedientes que la persona puede ver, con su estado de lectura.
    La búsqueda usa el buscador inteligente (sin acentos, frases como «de Luis» o «amparo»)."""
    from modules import buscador
    from modules.asuntos import TIPOS_DERECHO, _FILTRO_PROPIOS, _solo_lo_suyo, _usuarios
    where, params = ["1"], []
    if _solo_lo_suyo():
        uid = session.get("user_id")
        where.append(_FILTRO_PROPIOS)
        params += [uid, uid, uid]
    w = " AND ".join(where)
    # La tabla del lector nace con el primer archivo; un despacho nuevo aún no la tiene
    lector.indice.asegurar_esquema(tdb)
    filas = [dict(r, origen="documento", tabla="documentos") for r in tdb.execute(
        f"""SELECT d.id, d.nombre, d.archivo_url, d.archivo_nombre, d.tamanio, d.fecha_subida, d.subido_por,
                   d.tipo_documento AS clase, d.tipo_documento, d.tipo_estado, d.fecha_documento,
                   a.id AS asunto_id, a.titulo AS asunto, a.numero_expediente,
                   a.tipo_asunto, a.abogado_responsable, a.estado AS asunto_estado,
                   l.estado AS lectura, l.error AS lectura_error, l.paginas
            FROM documentos d
            JOIN asuntos a ON a.id = d.asunto_id
            LEFT JOIN lector_documentos l ON l.archivo_url = d.archivo_url
            WHERE {w}""", params).fetchall()]
    filas += [dict(r, origen="resolucion", tabla="asunto_sentencias", archivo_nombre=None, tamanio=None,
                   lectura=None, lectura_error=None, paginas=None)
              for r in tdb.execute(
        f"""SELECT s.id, s.titulo AS nombre, s.archivo_url, COALESCE(s.fecha_emision, s.fecha_alta) AS fecha_subida,
                   s.creado_por AS subido_por, s.tipo AS clase, s.contenido, s.tipo_documento, s.tipo_estado,
                   NULL AS fecha_documento, a.id AS asunto_id, a.titulo AS asunto,
                   a.numero_expediente, a.tipo_asunto, a.abogado_responsable, a.estado AS asunto_estado
            FROM asunto_sentencias s JOIN asuntos a ON a.id = s.asunto_id
            WHERE {w}""", params).fetchall()]

    q = (request.args.get("q") or "").strip()
    chips, sugerencia = [], None
    usuarios = _usuarios()
    if q:
        inter = buscador.interpretar(q, usuarios)
        f = inter["filtros"]
        chips = inter["chips"]
        if f.get("materia"):
            filas = [d for d in filas if d["tipo_asunto"] == f["materia"]]
        uid = buscador._uid_responsable(f)
        if uid:
            filas = [d for d in filas if d["abogado_responsable"] == uid]
        for d in filas:
            d["_pajar"] = buscador.normalizar(" ".join(str(v or "") for v in (
                d["nombre"], d["archivo_nombre"], tipos.etiqueta(d["tipo_documento"]), d["asunto"], d["numero_expediente"],
                TIPOS_DERECHO.get(d["tipo_asunto"] or "otro", {}).get("label"), (d.get("contenido") or "")[:400])))
        filas, corregido = buscador._filtrar(filas, inter["tokens"])
        sugerencia = " ".join(corregido) if corregido else None

    filas.sort(key=lambda d: str(d["fecha_subida"] or ""), reverse=True)
    docs, conteo = [], {"listo": 0, "procesando": 0, "revision": 0, "error": 0, "guardado": 0}
    carpetas: dict[int, dict] = {}
    for d in filas:
        info = ESTADOS_LECTURA.get(d["lectura"] or "")
        if d["origen"] == "resolucion":
            d["estado"], d["grupo"] = ("Guardado", "guardado") if d["archivo_url"] else ("Capturada", "guardado")
        else:
            d["estado"] = info["label"] if info else ("Guardado" if d["archivo_url"] else "Sin archivo")
            d["grupo"] = info["grupo"] if info else "guardado"
        nombre_arch = d["archivo_nombre"] or d["archivo_url"] or ""
        d["ext"] = nombre_arch.rsplit(".", 1)[-1].lower()[:4] if "." in nombre_arch else ""
        d["tipo_txt"] = "Resolución" if d["origen"] == "resolucion" else (d["ext"].upper() or "Documento")
        d["autor"] = (usuarios.get(d["subido_por"]) or {}).get("corto", "")
        d["dia"] = _dia(d.get("fecha_documento") or d["fecha_subida"])
        _con_tipo(d)
        conteo[d["grupo"]] += 1
        c = carpetas.setdefault(d["asunto_id"], {"id": d["asunto_id"], "titulo": d["asunto"], "n": 0,
                                                 "numero": d["numero_expediente"] or f"#{d['asunto_id']}",
                                                 "abierto": d["asunto_estado"] not in ("cerrado", "completado", "archivado")})
        c["n"] += 1
        docs.append(d)
    asunto_f = request.args.get("asunto", type=int)
    if asunto_f:
        docs = [d for d in docs if d["asunto_id"] == asunto_f]
    estado_f = request.args.get("estado", "")
    if estado_f in conteo:
        docs = [d for d in docs if d["grupo"] == estado_f]
    carpetas_l = sorted(carpetas.values(), key=lambda c: (not c["abierto"], -c["n"]))
    return docs, conteo, {"chips": chips, "sugerencia": sugerencia, "carpetas": carpetas_l, "asunto": asunto_f}


# ── Bóveda: «¿Qué es?» — tipo de cada documento y filtros por grupo/tipo ────────

SECCIONES = {
    "expedientes": {"label": "Expedientes",
                    "ayuda": "Escritos, acuerdos, sentencias y pruebas de cada expediente. Para agregar uno: "
                             "«Subir documento» y elige el expediente (o súbelo desde el propio expediente)."},
    "personas":    {"label": "Personas",
                    "ayuda": "Identificaciones, comprobantes, actas y contratos de tus clientes que no son de un "
                             "expediente. Para agregar uno: «Subir documento» y elige la persona."},
    "estudio":     {"label": "Archivo del despacho",
                    "ayuda": "Doctrina, jurisprudencia, leyes y papeles internos que no son de ningún expediente. "
                             "Para agregar: «Subir documento» → «Archivo del despacho»."},
    "formatos":    {"label": "Machotes y formatos",
                    "ayuda": "Modelos de escritos y contratos que se llenan solos con los datos de un expediente. "
                             "Para agregar: «Subir documento» → «Machotes y formatos»."},
}
# «Acuerdos y audiencias» ya no es pestaña: los acuerdos y las actas de audiencia son tipos de documento
# (filtro «Del juzgado o autoridad»). Los registros viejos siguen en ?seccion=acuerdos.


def _con_tipo(d: dict) -> dict:
    """Agrega a la fila su tipo legible: etiqueta, grupo y si solo es sugerido."""
    clave = tipos.valido(d.get("tipo_documento"))
    d["tipo"] = clave
    d["tipo_label"] = tipos.etiqueta(clave)
    d["tipo_grupo"] = tipos.grupo_de(clave)
    d["tipo_sugerido"] = bool(clave) and d.get("tipo_estado") != tipos.CONFIRMADO
    return d


def _filtro_tipos(filas: list[dict]) -> tuple[list[dict], dict]:
    """Chips por grupo (con conteo), que se abren en sus tipos; más «Sin clasificar» y «Por confirmar».
    Se combinan con la sección, el filtro por nombre y la búsqueda dentro de los documentos."""
    grupo_f = request.args.get("grupo", "")
    tipo_f = tipos.valido(request.args.get("tipo"))
    if tipo_f:
        grupo_f = tipos.grupo_de(tipo_f)
    por_grupo, por_tipo, sin, confirmar = {}, {}, 0, 0
    for d in filas:
        if d.get("tipo"):
            por_grupo[d["tipo_grupo"]] = por_grupo.get(d["tipo_grupo"], 0) + 1
            por_tipo[d["tipo"]] = por_tipo.get(d["tipo"], 0) + 1
            confirmar += d["tipo_sugerido"]
        else:
            sin += 1
    grupos = []
    for gr in tipos.por_grupo():
        n = por_grupo.get(gr["clave"], 0)
        if not n and grupo_f != gr["clave"]:
            continue
        grupos.append({"clave": gr["clave"], "etiqueta": gr["etiqueta"], "n": n, "on": grupo_f == gr["clave"],
                       "tipos": [{"clave": t["clave"], "etiqueta": t["etiqueta"], "n": por_tipo.get(t["clave"], 0),
                                  "on": tipo_f == t["clave"]}
                                 for t in gr["tipos"] if por_tipo.get(t["clave"]) or tipo_f == t["clave"]]})
    if tipo_f:
        filas = [d for d in filas if d.get("tipo") == tipo_f]
    elif grupo_f == "sin":
        filas = [d for d in filas if not d.get("tipo")]
    elif grupo_f == "confirmar":
        filas = [d for d in filas if d.get("tipo_sugerido")]
    elif grupo_f in tipos.GRUPO_LABEL:
        filas = [d for d in filas if d.get("tipo_grupo") == grupo_f]
    else:
        grupo_f = ""
    return filas, {"grupos": grupos, "sin": sin, "confirmar": confirmar, "grupo": grupo_f, "tipo": tipo_f,
                   "activo": bool(grupo_f or tipo_f)}


def _documentos_de_personas(tdb) -> list[dict]:
    """Documentos de las personas (clientes activos), con su estado de lectura y su tipo."""
    from modules.asuntos import _usuarios
    lector.indice.asegurar_esquema(tdb)
    filas = [dict(r, origen="cliente", tabla="cliente_documentos") for r in tdb.execute(
        """SELECT cd.id, cd.nombre, cd.archivo_url, cd.archivo_nombre, cd.tamanio, cd.fecha_subida, cd.subido_por,
                  cd.fecha_documento, cd.tipo_documento, cd.tipo_estado, c.id AS cliente_id,
                  COALESCE(NULLIF(c.razon_social, ''), TRIM(COALESCE(c.nombre,'') || ' ' || COALESCE(c.apellidos,''))) AS cliente,
                  l.estado AS lectura, l.error AS lectura_error, l.paginas
           FROM cliente_documentos cd JOIN clientes c ON c.id = cd.cliente_id AND c.activo = 1
           LEFT JOIN lector_documentos l ON l.archivo_url = cd.archivo_url
           ORDER BY COALESCE(cd.fecha_documento, cd.fecha_subida) DESC""").fetchall()]
    q = _plano(request.args.get("q") or "")
    if q:
        filas = [d for d in filas if all(t in _plano(" ".join(str(v or "") for v in (
            d["nombre"], d["archivo_nombre"], d["cliente"], tipos.etiqueta(d["tipo_documento"])))) for t in q.split())]
    usuarios = _usuarios()
    for d in filas:
        info = ESTADOS_LECTURA.get(d["lectura"] or "")
        d["estado"] = info["label"] if info else "Guardado"
        d["grupo"] = info["grupo"] if info else "guardado"
        nombre_arch = d["archivo_nombre"] or d["archivo_url"] or ""
        d["ext"] = nombre_arch.rsplit(".", 1)[-1].lower()[:4] if "." in nombre_arch else ""
        d["autor"] = (usuarios.get(d["subido_por"]) or {}).get("corto", "")
        d["dia"] = _dia(d["fecha_documento"] or d["fecha_subida"])
        _con_tipo(d)
    return filas


def _puede_clasificar(tdb, tabla: str, fila_id: int) -> bool:
    """Misma regla que para verlos: lo de un expediente, solo quien puede ver ese expediente."""
    from modules.asuntos import puede_ver
    if tabla not in tipos.TABLAS:
        return False
    if tabla == "cliente_documentos":
        return bool(tdb.execute("SELECT 1 FROM cliente_documentos cd JOIN clientes c ON c.id = cd.cliente_id "
                                "WHERE cd.id=? AND c.activo=1", (fila_id,)).fetchone())
    fila = tdb.execute(f"SELECT asunto_id FROM {tabla} WHERE id=?", (fila_id,)).fetchone()
    if not fila:
        return False
    if fila["asunto_id"] is None:
        return tabla == "biblioteca_archivos"
    return puede_ver(tdb, fila["asunto_id"])


@biblioteca_bp.route("/tipo", methods=["POST"])
@login_required
@tenant_required
def tipo_guardar():
    """Clasificar un documento desde su tarjeta: «¿Qué es?». Queda confirmado por la persona."""
    data = request.get_json(silent=True) or request.form
    tabla = data.get("tabla", "")
    fila_id = _safe_int(data.get("id"))
    clave = tipos.valido(data.get("tipo"))
    if not clave:
        return jsonify(ok=False, error="Elige qué es el documento"), 400
    if not fila_id or not _puede_clasificar(g.tenant_db, tabla, fila_id):
        return jsonify(ok=False, error="No encontré ese documento"), 404
    tipos.guardar(g.tenant_db, tabla, fila_id, clave, confirmado=True)
    g.tenant_db.commit()
    _log("clasificar_documento", fila_id, {"tabla": tabla, "tipo": clave})
    return jsonify(ok=True, tipo=clave, etiqueta=tipos.etiqueta(clave),
                   grupo=tipos.GRUPO_LABEL.get(tipos.grupo_de(clave), ""))


@biblioteca_bp.route("/tipo/sugerir")
@login_required
@tenant_required
def tipo_sugerir():
    """Sugerencia por el nombre del archivo (la misma que usa el servidor al guardar)."""
    clave = tipos.sugerir((request.args.get("nombre") or "")[:200])
    return jsonify(ok=True, tipo=clave, etiqueta=tipos.etiqueta(clave))


# ── Bóveda: búsqueda dentro del contenido de los documentos ──────────────────

ORIGEN_TXT = {"documento": "Documento del expediente", "sentencia": "Resolución", "cliente": "Documento de la persona",
              "biblioteca": "Estudio", "plantilla": "Formato"}
_VACIAS = {"de", "la", "el", "los", "las", "del", "que", "por", "con", "una", "uno", "para", "en", "se", "su", "sus",
           "al", "lo", "un", "y", "o", "a", "es", "no", "como", "sobre", "entre", "este", "esta"}
MAX_PASAJES = 4
_TABLA_DE_ORIGEN = {"documento": "documentos", "sentencia": "asunto_sentencias", "cliente": "cliente_documentos",
                    "biblioteca": "biblioteca_archivos"}


def _plano(t: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", t.lower()) if not unicodedata.combining(c))


def _terminos(consulta: str) -> list[str]:
    return [t for t in dict.fromkeys(_plano(w) for w in re.findall(r"\w+", consulta))
            if t not in _VACIAS and (len(t) > 2 or t.isdigit())]


def _coincide(palabra: str, terminos: list[str]) -> bool:
    w = _plano(palabra)
    return any(w == t or (len(t) >= 4 and w.startswith(t[:5])) for t in terminos)


def _resaltar(texto: str, terminos: list[str]) -> Markup:
    """Texto escapado con las palabras buscadas dentro de <mark> (sin acentos ni mayúsculas, por raíz)."""
    partes, pos = [], 0
    for m in re.finditer(r"\w+", texto):
        if _coincide(m.group(), terminos):
            partes.append(escape(texto[pos:m.start()]))
            partes.append(Markup("<mark>%s</mark>") % m.group())
            pos = m.end()
    partes.append(escape(texto[pos:]))
    return Markup("").join(partes)


def _dia(v) -> str:
    try:
        f = date.fromisoformat(str(v or "")[:10])
    except ValueError:
        return ""
    return f"{f.day} {_MESES_CORTOS[f.month - 1]} {f.year}"


def _buscar_en_contenido(tdb, consulta: str, visibles: dict[int, dict]) -> list[dict]:
    """Resultados agrupados por documento: cada uno con sus pasajes, página y palabras resaltadas.
    Solo entran los documentos que la persona puede ver (lector.documentos_visibles)."""
    permitidos = {i for i, d in visibles.items() if d["estado"] == "listo"}
    if not consulta or not permitidos:
        return []
    hallazgos = lector.indice.buscar(current_app._get_current_object(), g.tenant_slug, tdb, consulta,
                                     limite=40, documentos=permitidos)
    terminos = _terminos(consulta)
    grupos: dict[int, dict] = {}
    for h in hallazgos:
        d = visibles.get(h["doc_id"])
        if not d:
            continue
        grupo = grupos.get(h["doc_id"])
        if grupo is None:
            url = d["archivo_url"] or ""
            ext = url.rsplit(".", 1)[-1].lower()[:4] if "." in url else ""
            grupo = grupos[h["doc_id"]] = {
                "id": h["doc_id"], "nombre": d["nombre"] or h["documento"] or "Documento sin nombre",
                "url": url if boveda.es_de_boveda(url) else None, "ext": ext, "pdf": ext == "pdf",
                "origen": ORIGEN_TXT.get(d["origen"], "Documento"), "dia": _dia(d["fecha"]),
                "asunto_id": d["asunto_id"], "asunto": d["asunto"], "numero": d["numero_expediente"],
                "cliente_id": d["cliente_id"], "cliente": (d["cliente"] or "").strip() or None,
                "paginas": d["paginas"], "pasajes": [], "seccion": "formatos" if d["origen"] == "plantilla" else
                ("estudio" if d["origen"] == "biblioteca" else None),
                "tabla": _TABLA_DE_ORIGEN.get(d["origen"]), "fila_id": d["fila_id"],
                "tipo_documento": d["tipo_documento"], "tipo_estado": d["tipo_estado"],
            }
            _con_tipo(grupo)
        grupo["pasajes"].append({"pagina": h["pagina"] or 1, "orden": h["fragmento_id"],
                                 "html": _resaltar(h["resumen"], terminos)})
    salida = list(grupos.values())
    for grupo in salida:
        grupo["pasajes"].sort(key=lambda x: (x["pagina"], x["orden"]))
        grupo["n"] = len(grupo["pasajes"])
    return salida


@biblioteca_bp.route("/")
@login_required
@tenant_required
def index():
    from modules.asuntos import _FILTRO_PROPIOS, _solo_lo_suyo
    tdb = g.tenant_db
    seccion = request.args.get("seccion", "expedientes")
    if seccion not in (*SECCIONES, "acuerdos"):
        seccion = "expedientes"
    # Documentos que nunca se han clasificado: se sugiere su tipo una sola vez (queda «sugerido»)
    tipos.clasificar_pendientes(tdb, g.tenant_slug)

    # ── Búsqueda dentro de los documentos y estado de lectura (solo lo que la persona puede ver) ──
    texto = (request.args.get("texto") or "").strip()[:300]
    visibles = lector.documentos_visibles(tdb)
    lectura = lector.conteo_lectura(visibles)
    lectura["reintentar"] = lector.puede_reintentar() and lectura["reintentables"] > 0
    filtro = None
    en_contenido = _buscar_en_contenido(tdb, texto, visibles) if texto else []
    if texto:
        en_contenido, filtro = _filtro_tipos(en_contenido)
    docs_exp, conteo_estados, busqueda = (_documentos_de_expedientes(tdb) if seccion == "expedientes" and not texto
                                          else ([], {}, {"chips": [], "sugerencia": None, "carpetas": [], "asunto": None}))
    if seccion == "expedientes" and not texto:
        docs_exp, filtro = _filtro_tipos(docs_exp)
    docs_personas = []
    if seccion == "personas" and not texto:
        docs_personas, filtro = _filtro_tipos(_documentos_de_personas(tdb))

    # ── Estudio data ──
    carpetas = tdb.execute(
        "SELECT c.*, (SELECT COUNT(*) FROM biblioteca_archivos a WHERE a.carpeta_id = c.id) AS n "
        "FROM biblioteca_carpetas c ORDER BY c.nombre"
    ).fetchall()
    carpetas = [dict(c) for c in carpetas]

    carpeta_id = _safe_int(request.args.get("carpeta"))
    q = request.args.get("q", "").strip()
    tipo_f = request.args.get("ext", "")

    where, params = ["1"], []
    if _solo_lo_suyo():
        # Un archivo vinculado a un expediente solo lo ve quien puede ver ese expediente
        uid = session.get("user_id")
        where.append(f"(a.asunto_id IS NULL OR a.asunto_id IN (SELECT a.id FROM asuntos a WHERE {_FILTRO_PROPIOS}))")
        params += [uid, uid, uid]
    if carpeta_id:
        where.append("a.carpeta_id=?")
        params.append(carpeta_id)
    if q:
        where.append("(a.nombre LIKE ? OR a.nombre_original LIKE ?)")
        params += [f"%{q}%", f"%{q}%"]
    if tipo_f:
        where.append("a.extension=?")
        params.append(tipo_f)

    archivos = tdb.execute(
        f"SELECT a.*, c.nombre as carpeta_nombre "
        f"FROM biblioteca_archivos a LEFT JOIN biblioteca_carpetas c ON a.carpeta_id=c.id "
        f"WHERE {' AND '.join(where)} ORDER BY a.fecha_alta DESC LIMIT 200",
        params,
    ).fetchall()
    archivos = [_con_tipo(dict(a)) for a in archivos]
    for a in archivos:
        a["dia"] = _dia(a.get("fecha_documento") or a.get("fecha_alta"))
    if seccion == "estudio" and not texto:
        archivos, filtro = _filtro_tipos(archivos)

    espacio_usado = _espacio_usado(tdb)
    espacio_limite = _plan_limit_bytes()
    total_archivos = tdb.execute("SELECT COUNT(*) FROM biblioteca_archivos").fetchone()[0]

    # ── Formatos data ──
    cat_f = request.args.get("categoria", "")
    derecho_f = request.args.get("derecho", "")
    q_pl = request.args.get("q_pl", "").strip()

    pw, pp = ["1"], []
    if cat_f:
        pw.append("p.categoria=?"); pp.append(cat_f)
    if derecho_f:
        pw.append("p.tipo_derecho=?"); pp.append(derecho_f)
    if q_pl:
        pw.append("(p.nombre LIKE ? OR p.descripcion LIKE ?)")
        pp += [f"%{q_pl}%", f"%{q_pl}%"]

    plantillas = tdb.execute(
        f"SELECT * FROM biblioteca_plantillas p WHERE {' AND '.join(pw)} ORDER BY p.nombre",
        pp,
    ).fetchall()
    plantillas = [dict(p) for p in plantillas]

    # ── Acuerdos y audiencias (registros anteriores; ya no es pestaña) ──
    n_acuerdos = tdb.execute("SELECT COUNT(*) FROM biblioteca_acuerdos").fetchone()[0]
    tab_ac = request.args.get("tab_ac", "audiencias")
    acuerdos = tdb.execute(
        "SELECT ba.*, a.titulo as asunto_titulo FROM biblioteca_acuerdos ba "
        "LEFT JOIN asuntos a ON ba.asunto_id=a.id "
        "WHERE ba.tipo=? ORDER BY ba.fecha DESC LIMIT 100",
        ("audiencia" if tab_ac == "audiencias" else "acuerdo",),
    ).fetchall()
    acuerdos = [dict(a) for a in acuerdos]

    # ── Listas para «¿Dónde va?» (solo los expedientes que la persona puede ver) ──
    alc, alc_p = ("1", [])
    if _solo_lo_suyo():
        uid = session.get("user_id")
        alc, alc_p = _FILTRO_PROPIOS, [uid, uid, uid]
    asuntos = tdb.execute(
        f"SELECT a.id, a.titulo, a.numero_expediente, a.estado FROM asuntos a WHERE a.estado != 'archivado' AND {alc} "
        "ORDER BY CASE WHEN a.estado IN ('cerrado','completado') THEN 1 ELSE 0 END, a.titulo", alc_p
    ).fetchall()
    clientes = tdb.execute(
        "SELECT id, nombre, apellidos, razon_social FROM clientes WHERE activo=1 ORDER BY nombre LIMIT 1000"
    ).fetchall()
    destinos = {
        "expedientes": [{"id": a["id"], "t": f"{a['numero_expediente'] or '#' + str(a['id'])} · {a['titulo']}"} for a in asuntos],
        "personas": [{"id": c["id"], "t": (c["razon_social"] or f"{c['nombre'] or ''} {c['apellidos'] or ''}").strip()}
                     for c in clientes],
        "carpetas": [{"id": c["id"], "t": c["nombre"]} for c in carpetas],
    }

    resp = make_response(render_template(
        "biblioteca/index.html",
        texto=texto, en_contenido=en_contenido, lectura=lectura,
        max_pasajes=MAX_PASAJES,
        seccion=seccion, secciones=SECCIONES, filtro=filtro, docs_personas=docs_personas,
        tipos_grupos=tipos.por_grupo(), destinos=destinos, n_acuerdos=n_acuerdos,
        puede_formatos=session.get("role") in ("super_admin", "admin", "abogado"),
        docs_exp=docs_exp, conteo_estados=conteo_estados, estado_f=request.args.get("estado", ""), busqueda=busqueda,
        estados_lectura=ESTADOS_LECTURA,
        # Estudio
        carpetas=carpetas, archivos=archivos,
        carpeta_id=carpeta_id, q=q, tipo_f=tipo_f,
        espacio_usado=espacio_usado, espacio_limite=espacio_limite,
        total_archivos=total_archivos,
        # Formatos
        plantillas=plantillas, categorias=CATEGORIAS_PLANTILLA,
        tipos_derecho=TIPOS_DERECHO,
        cat_f=cat_f, derecho_f=derecho_f, q_pl=q_pl,
        # Acuerdos
        acuerdos=acuerdos, tab_ac=tab_ac,
        # Shared
        asuntos=asuntos, clientes=clientes,
        abogados=_get_abogados(),
    ))
    if texto:
        # Los pasajes son texto de los documentos: que el navegador no los guarde
        resp.headers["Cache-Control"] = "private, no-store"
    return resp


# ═════════════════════════════════════════════════════════════════════════════
#  ESTUDIO — Carpetas
# ═════════════════════════════════════════════════════════════════════════════

@biblioteca_bp.route("/carpeta/crear", methods=["POST"])
@login_required
@tenant_required
def carpeta_crear():
    nombre = (request.form.get("nombre") or "").strip()
    parent_id = _safe_int(request.form.get("parent_id"))
    if not nombre:
        return jsonify(ok=False, error="Nombre requerido"), 400
    cur = g.tenant_db.execute(
        "INSERT INTO biblioteca_carpetas (nombre, parent_id, creado_por) VALUES (?,?,?)",
        (nombre, parent_id, session.get("user_id")),
    )
    g.tenant_db.commit()
    _log("crear_carpeta", cur.lastrowid, {"nombre": nombre})
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return jsonify(ok=True, id=cur.lastrowid, nombre=nombre)
    return redirect(url_for("biblioteca.index", seccion="estudio"))


@biblioteca_bp.route("/carpeta/<int:id>/renombrar", methods=["POST"])
@login_required
@tenant_required
def carpeta_renombrar(id):
    nombre = (request.form.get("nombre") or request.json.get("nombre", "") if request.is_json else "").strip()
    if not nombre:
        return jsonify(ok=False, error="Nombre requerido"), 400
    g.tenant_db.execute("UPDATE biblioteca_carpetas SET nombre=? WHERE id=?", (nombre, id))
    g.tenant_db.commit()
    return jsonify(ok=True)


@biblioteca_bp.route("/carpeta/<int:id>/eliminar", methods=["POST"])
@login_required
@tenant_required
@role_required("super_admin", "admin")
def carpeta_eliminar(id):
    g.tenant_db.execute("DELETE FROM biblioteca_carpetas WHERE id=?", (id,))
    g.tenant_db.execute("UPDATE biblioteca_archivos SET carpeta_id=NULL WHERE carpeta_id=?", (id,))
    g.tenant_db.commit()
    _log("eliminar_carpeta", id)
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return jsonify(ok=True)
    return redirect(url_for("biblioteca.index", seccion="estudio"))


# ═════════════════════════════════════════════════════════════════════════════
#  ESTUDIO — Archivos
# ═════════════════════════════════════════════════════════════════════════════

@biblioteca_bp.route("/upload", methods=["POST"])
@login_required
@tenant_required
def upload():
    from modules.asuntos import puede_ver
    tdb = g.tenant_db
    tipos.asegurar(tdb)
    carpeta_id = _safe_int(request.form.get("carpeta_id"))
    asunto_id = _safe_int(request.form.get("asunto_id"))
    cliente_id = _safe_int(request.form.get("cliente_id"))
    if asunto_id and not puede_ver(tdb, asunto_id):
        return jsonify(ok=False, error="No encontré ese expediente"), 404

    # Check space
    usado = _espacio_usado(tdb)
    limite = _plan_limit_bytes()

    uploaded = []
    files = request.files.getlist("archivos")
    # Desde «Subir documento» llega un solo archivo y puede traer un nombre propio
    nombre_propio = (request.form.get("nombre") or "").strip()[:200] if len(files) == 1 else ""
    for f in files:
        if usado >= limite:
            break
        info = _save_file(f, "estudio")
        if not info:
            continue
        cur = tdb.execute(
            """INSERT INTO biblioteca_archivos
               (carpeta_id, nombre, nombre_original, archivo_url, tamanio,
                mime_type, extension, asunto_id, cliente_id, creado_por)
               VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (
                carpeta_id,
                nombre_propio or info["nombre_original"],
                info["nombre_original"],
                info["url"],
                info["tamanio"],
                info["mime_type"],
                info["extension"],
                asunto_id, cliente_id,
                session.get("user_id"),
            ),
        )
        usado += info["tamanio"]
        # «¿Qué es?»: lo que eligió la persona (confirmado) o, si no dijo, sugerido por el nombre
        tipos.al_subir(tdb, "biblioteca_archivos", cur.lastrowid, request.form.get("tipo_documento"),
                       info["nombre_original"], request.form.get("fecha_documento"))
        uploaded.append({"id": cur.lastrowid, "nombre": info["nombre_original"], "url": info["url"]})
    tdb.commit()
    for u in uploaded:
        lector.encolar(u.pop("url"), u["nombre"], "biblioteca", u["id"], asunto_id=asunto_id, cliente_id=cliente_id)
    _log("upload_archivos", None, {"count": len(uploaded)})
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return jsonify(ok=True, archivos=uploaded, count=len(uploaded))
    return redirect(url_for("biblioteca.index", seccion="estudio", carpeta=carpeta_id or ""))


@biblioteca_bp.route("/archivo/<int:id>/eliminar", methods=["POST"])
@login_required
@tenant_required
def archivo_eliminar(id):
    fila = g.tenant_db.execute("SELECT archivo_url FROM biblioteca_archivos WHERE id=?", (id,)).fetchone()
    g.tenant_db.execute("DELETE FROM biblioteca_archivos WHERE id=?", (id,))
    g.tenant_db.commit()
    if fila:
        boveda.eliminar(fila["archivo_url"])
    _log("eliminar_archivo", id)
    return jsonify(ok=True)


@biblioteca_bp.route("/archivo/<int:id>/compartir", methods=["POST"])
@login_required
@tenant_required
def archivo_compartir(id):
    data = request.json or request.form
    horas = _safe_int(data.get("horas"), 24)
    if horas not in (24, 48, 72):
        horas = 24
    token = secrets.token_urlsafe(32)
    expires = datetime.now() + timedelta(hours=horas)
    g.tenant_db.execute(
        """INSERT INTO biblioteca_compartidos (tipo, entidad_id, token, expires_at, creado_por)
           VALUES ('archivo', ?, ?, ?, ?)""",
        (id, token, expires.isoformat(), session.get("user_id")),
    )
    g.tenant_db.commit()
    link = url_for("biblioteca.compartido_ver", token=token, _external=True)
    return jsonify(ok=True, link=link, expires=expires.isoformat())


@biblioteca_bp.route("/archivo/<int:id>/vincular", methods=["POST"])
@login_required
@tenant_required
def archivo_vincular(id):
    data = request.json or request.form
    asunto_id = _safe_int(data.get("asunto_id"))
    cliente_id = _safe_int(data.get("cliente_id"))
    sets, vals = [], []
    if asunto_id is not None:
        sets.append("asunto_id=?"); vals.append(asunto_id)
    if cliente_id is not None:
        sets.append("cliente_id=?"); vals.append(cliente_id)
    if sets:
        vals.append(id)
        g.tenant_db.execute(f"UPDATE biblioteca_archivos SET {','.join(sets)} WHERE id=?", vals)
        g.tenant_db.commit()
    return jsonify(ok=True)


# ═════════════════════════════════════════════════════════════════════════════
#  FORMATOS — Plantillas
# ═════════════════════════════════════════════════════════════════════════════

@biblioteca_bp.route("/plantilla/subir", methods=["POST"])
@login_required
@tenant_required
@role_required("super_admin", "admin", "abogado")
def plantilla_subir():
    f = request.files.get("archivo")
    nombre = (request.form.get("nombre") or "").strip()
    if not nombre:
        return jsonify(ok=False, error="Nombre requerido"), 400
    info = _save_file(f, "plantillas")
    if not info:
        return jsonify(ok=False, error="Archivo inválido o no soportado"), 400
    cur = g.tenant_db.execute(
        """INSERT INTO biblioteca_plantillas
           (nombre, categoria, descripcion, tipo_derecho,
            archivo_url, archivo_nombre, tamanio, es_custom, creado_por)
           VALUES (?,?,?,?,?,?,?,1,?)""",
        (
            nombre,
            request.form.get("categoria", "Otro"),
            (request.form.get("descripcion") or "").strip() or None,
            request.form.get("tipo_derecho", "Otro"),
            info["url"], info["nombre_original"], info["tamanio"],
            session.get("user_id"),
        ),
    )
    g.tenant_db.commit()
    lector.encolar(info["url"], nombre, "plantilla", cur.lastrowid)
    _log("subir_plantilla", cur.lastrowid, {"nombre": nombre})
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return jsonify(ok=True, id=cur.lastrowid)
    return redirect(url_for("biblioteca.index", seccion="formatos"))


@biblioteca_bp.route("/plantilla/<int:id>/descargar")
@login_required
@tenant_required
def plantilla_descargar(id):
    p = g.tenant_db.execute("SELECT * FROM biblioteca_plantillas WHERE id=?", (id,)).fetchone()
    if not p:
        abort(404)
    nombre = p["archivo_nombre"] or f"{p['nombre']}.docx"
    if boveda.es_de_boveda(p["archivo_url"]):
        return boveda.respuesta(p["archivo_url"], nombre, descargar=True)
    from config import Config
    filepath = os.path.join(Config.UPLOAD_FOLDER, p["archivo_url"].lstrip("/uploads/").lstrip("/"))
    if not os.path.exists(filepath):
        # Try alternate path
        filepath = os.path.join(os.path.dirname(Config.UPLOAD_FOLDER), p["archivo_url"].lstrip("/"))
    if not os.path.exists(filepath):
        abort(404)
    return send_file(filepath, as_attachment=True, download_name=nombre)


@biblioteca_bp.route("/plantilla/<int:id>/eliminar", methods=["POST"])
@login_required
@tenant_required
@role_required("super_admin", "admin")
def plantilla_eliminar(id):
    fila = g.tenant_db.execute("SELECT archivo_url FROM biblioteca_plantillas WHERE id=?", (id,)).fetchone()
    g.tenant_db.execute("DELETE FROM biblioteca_plantillas WHERE id=?", (id,))
    g.tenant_db.commit()
    if fila:
        boveda.eliminar(fila["archivo_url"])
    return jsonify(ok=True)


# ═════════════════════════════════════════════════════════════════════════════
#  ACUERDOS Y AUDIENCIAS
# ═════════════════════════════════════════════════════════════════════════════

@biblioteca_bp.route("/acuerdo/crear", methods=["POST"])
@login_required
@tenant_required
def acuerdo_crear():
    data = request.json or request.form
    tipo = data.get("tipo", "acuerdo")
    if tipo not in ("acuerdo", "audiencia"):
        tipo = "acuerdo"
    cur = g.tenant_db.execute(
        """INSERT INTO biblioteca_acuerdos
           (tipo, neun, numero_asunto, organo, circuito, autoridad,
            fecha, descripcion, asunto_id, creado_por)
           VALUES (?,?,?,?,?,?,?,?,?,?)""",
        (
            tipo,
            (data.get("neun") or "").strip() or None,
            (data.get("numero_asunto") or "").strip() or None,
            (data.get("organo") or "").strip() or None,
            (data.get("circuito") or "").strip() or None,
            (data.get("autoridad") or "").strip() or None,
            data.get("fecha") or None,
            (data.get("descripcion") or "").strip() or None,
            _safe_int(data.get("asunto_id")),
            session.get("user_id"),
        ),
    )
    g.tenant_db.commit()
    _log("crear_acuerdo", cur.lastrowid)
    return jsonify(ok=True, id=cur.lastrowid)


@biblioteca_bp.route("/acuerdo/<int:id>/vincular", methods=["POST"])
@login_required
@tenant_required
def acuerdo_vincular(id):
    data = request.json or request.form
    asunto_id = _safe_int(data.get("asunto_id"))
    g.tenant_db.execute(
        "UPDATE biblioteca_acuerdos SET asunto_id=? WHERE id=?", (asunto_id, id)
    )
    g.tenant_db.commit()
    return jsonify(ok=True)


@biblioteca_bp.route("/acuerdo/<int:id>/eliminar", methods=["POST"])
@login_required
@tenant_required
def acuerdo_eliminar(id):
    g.tenant_db.execute("DELETE FROM biblioteca_acuerdos WHERE id=?", (id,))
    g.tenant_db.commit()
    return jsonify(ok=True)


# ═════════════════════════════════════════════════════════════════════════════
#  EDITOR DE DOCUMENTOS
# ═════════════════════════════════════════════════════════════════════════════

@biblioteca_bp.route("/editor")
@biblioteca_bp.route("/editor/<int:doc_id>")
@login_required
@tenant_required
def editor(doc_id=None):
    tdb = g.tenant_db
    documento = None
    versiones = []

    if doc_id:
        documento = tdb.execute("SELECT * FROM biblioteca_documentos WHERE id=?", (doc_id,)).fetchone()
        if documento:
            documento = dict(documento)
            versiones = tdb.execute(
                "SELECT id, version, creado_por, fecha FROM biblioteca_doc_versiones "
                "WHERE documento_id=? ORDER BY version DESC LIMIT 10",
                (doc_id,),
            ).fetchall()
            versiones = [dict(v) for v in versiones]

    plantilla_id = _safe_int(request.args.get("plantilla"))
    plantilla_html = ""
    if plantilla_id and not documento:
        p = tdb.execute("SELECT * FROM biblioteca_plantillas WHERE id=?", (plantilla_id,)).fetchone()
        if p and boveda.es_de_boveda(p["archivo_url"]) and p["archivo_url"].endswith(".docx"):
            try:
                plantilla_html = machotes.a_html(boveda.leer(p["archivo_url"]))
            except Exception:
                plantilla_html = ""

    asuntos = tdb.execute(
        "SELECT id, titulo, numero_expediente, cliente_id FROM asuntos "
        "WHERE estado NOT IN ('archivado','cerrado') ORDER BY titulo"
    ).fetchall()
    clientes = tdb.execute(
        "SELECT id, nombre, apellidos, razon_social FROM clientes WHERE activo=1 ORDER BY nombre LIMIT 300"
    ).fetchall()

    plantillas = tdb.execute(
        "SELECT id, nombre, categoria FROM biblioteca_plantillas ORDER BY nombre"
    ).fetchall()

    # Recent documents for sidebar
    recientes = tdb.execute(
        "SELECT id, titulo, estado, fecha_mod FROM biblioteca_documentos ORDER BY fecha_mod DESC LIMIT 10"
    ).fetchall()

    return render_template(
        "biblioteca/editor.html",
        documento=documento,
        versiones=versiones,
        plantilla_html=plantilla_html,
        plantilla_id=plantilla_id,
        variables=VARIABLES_DINAMICAS,
        asuntos=[dict(a) for a in asuntos],
        clientes=[dict(c) for c in clientes],
        plantillas=[dict(p) for p in plantillas],
        recientes=[dict(r) for r in recientes],
    )


@biblioteca_bp.route("/editor/guardar", methods=["POST"])
@login_required
@tenant_required
def editor_guardar():
    tdb = g.tenant_db
    data = request.json or {}
    doc_id = _safe_int(data.get("id"))
    titulo = (data.get("titulo") or "Sin título").strip()
    contenido = data.get("contenido_html", "")
    asunto_id = _safe_int(data.get("asunto_id"))
    cliente_id = _safe_int(data.get("cliente_id"))
    plantilla_id = _safe_int(data.get("plantilla_id"))

    if doc_id:
        # Update existing
        tdb.execute(
            """UPDATE biblioteca_documentos
               SET titulo=?, contenido_html=?, asunto_id=?, cliente_id=?,
                   fecha_mod=CURRENT_TIMESTAMP
               WHERE id=?""",
            (titulo, contenido, asunto_id, cliente_id, doc_id),
        )
        # Save version
        last_ver = tdb.execute(
            "SELECT MAX(version) FROM biblioteca_doc_versiones WHERE documento_id=?",
            (doc_id,),
        ).fetchone()[0] or 0
        tdb.execute(
            """INSERT INTO biblioteca_doc_versiones
               (documento_id, contenido_html, version, creado_por)
               VALUES (?,?,?,?)""",
            (doc_id, contenido, last_ver + 1, session.get("user_id")),
        )
        # Keep only last 10 versions
        tdb.execute(
            """DELETE FROM biblioteca_doc_versiones
               WHERE documento_id=? AND id NOT IN (
                 SELECT id FROM biblioteca_doc_versiones
                 WHERE documento_id=? ORDER BY version DESC LIMIT 10
               )""",
            (doc_id, doc_id),
        )
        tdb.commit()
        return jsonify(ok=True, id=doc_id, version=last_ver + 1)
    else:
        # Create new
        cur = tdb.execute(
            """INSERT INTO biblioteca_documentos
               (titulo, contenido_html, plantilla_id, asunto_id, cliente_id, creado_por)
               VALUES (?,?,?,?,?,?)""",
            (titulo, contenido, plantilla_id, asunto_id, cliente_id, session.get("user_id")),
        )
        doc_id = cur.lastrowid
        # Version 1
        tdb.execute(
            "INSERT INTO biblioteca_doc_versiones (documento_id, contenido_html, version, creado_por) VALUES (?,?,1,?)",
            (doc_id, contenido, session.get("user_id")),
        )
        tdb.commit()
        _log("crear_documento", doc_id, {"titulo": titulo})
        return jsonify(ok=True, id=doc_id, version=1, nuevo=True)


@biblioteca_bp.route("/editor/<int:doc_id>/version/<int:ver_id>")
@login_required
@tenant_required
def editor_version(doc_id, ver_id):
    v = g.tenant_db.execute(
        "SELECT * FROM biblioteca_doc_versiones WHERE id=? AND documento_id=?",
        (ver_id, doc_id),
    ).fetchone()
    if not v:
        return jsonify(ok=False, error="Versión no encontrada"), 404
    return jsonify(ok=True, contenido_html=v["contenido_html"], version=v["version"], fecha=v["fecha"])


@biblioteca_bp.route("/editor/<int:doc_id>/variables", methods=["POST"])
@login_required
@tenant_required
def editor_resolver_variables(doc_id):
    """Replace variables in the document content and return resolved HTML."""
    data = request.json or {}
    html = data.get("contenido_html", "")
    asunto_id = _safe_int(data.get("asunto_id"))
    cliente_id = _safe_int(data.get("cliente_id"))
    resolved = _resolve_variables(html, asunto_id, cliente_id)
    return jsonify(ok=True, contenido_html=resolved)


@biblioteca_bp.route("/editor/exportar-pdf", methods=["POST"])
@login_required
@tenant_required
def editor_exportar_pdf():
    data = request.json or {}
    html = data.get("contenido_html", "")
    titulo = data.get("titulo", "Documento")
    asunto_id = _safe_int(data.get("asunto_id"))
    cliente_id = _safe_int(data.get("cliente_id"))

    # Resolve variables before export
    html = _resolve_variables(html, asunto_id, cliente_id)

    full_html = f"""<!DOCTYPE html>
<html><head>
<meta charset="utf-8">
<style>
  body {{ font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif; font-size: 12pt;
         line-height: 1.6; color: #1a1a1a; max-width: 700px; margin: 40px auto; padding: 0 20px; }}
  h1 {{ font-size: 18pt; color: #333; }} h2 {{ font-size: 15pt; color: #444; }}
  h3 {{ font-size: 13pt; color: #555; }} h4 {{ font-size: 12pt; color: #666; }}
  table {{ border-collapse: collapse; width: 100%; margin: 12px 0; }}
  th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
  th {{ background: #f5f5f5; font-weight: 600; }}
  blockquote {{ border-left: 3px solid #AD9060; padding-left: 12px; color: #555; margin: 12px 0; }}
  img {{ max-width: 100%; }}
</style>
</head><body>{html}</body></html>"""

    # Try WeasyPrint first, fallback to pdfkit, then reportlab
    buf = io.BytesIO()
    try:
        from weasyprint import HTML
        HTML(string=full_html).write_pdf(buf)
    except ImportError:
        try:
            import pdfkit
            pdf_bytes = pdfkit.from_string(full_html, False)
            buf.write(pdf_bytes)
        except ImportError:
            # Fallback: simple reportlab
            try:
                from reportlab.lib.pagesizes import A4
                from reportlab.platypus import SimpleDocTemplate, Paragraph
                from reportlab.lib.styles import getSampleStyleSheet
                doc = SimpleDocTemplate(buf, pagesize=A4)
                st = getSampleStyleSheet()
                # Strip HTML for reportlab
                import re
                text = re.sub(r'<[^>]+>', ' ', html).strip()
                story = [Paragraph(titulo, st["Heading1"]), Paragraph(text, st["Normal"])]
                doc.build(story)
            except ImportError:
                return jsonify(ok=False, error="No hay generador PDF instalado (weasyprint/pdfkit/reportlab)"), 500

    buf.seek(0)
    safe_name = titulo.replace(" ", "_")[:30]
    return send_file(buf, mimetype="application/pdf", as_attachment=True,
                     download_name=f"{safe_name}.pdf")


@biblioteca_bp.route("/editor/exportar-docx", methods=["POST"])
@login_required
@tenant_required
def editor_exportar_docx():
    data = request.json or {}
    html = data.get("contenido_html", "")
    titulo = data.get("titulo", "Documento")
    asunto_id = _safe_int(data.get("asunto_id"))
    cliente_id = _safe_int(data.get("cliente_id"))

    html = _resolve_variables(html, asunto_id, cliente_id)

    try:
        from docx import Document
        from docx.shared import Pt, Inches
        import re
    except ImportError:
        return jsonify(ok=False, error="python-docx no instalado"), 500

    doc = Document()
    doc.add_heading(titulo, level=1)

    # Simple HTML to DOCX conversion
    import re
    # Split by block elements
    blocks = re.split(r'<(?:p|div|h[1-4]|li|blockquote|tr)[^>]*>', html)
    for block in blocks:
        text = re.sub(r'<[^>]+>', '', block).strip()
        if text:
            doc.add_paragraph(text)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    safe_name = titulo.replace(" ", "_")[:30]
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                     as_attachment=True, download_name=f"{safe_name}.docx")


@biblioteca_bp.route("/editor/<int:doc_id>/compartir", methods=["POST"])
@login_required
@tenant_required
def editor_compartir(doc_id):
    token = secrets.token_urlsafe(32)
    horas = _safe_int((request.json or {}).get("horas"), 72)
    expires = datetime.now() + timedelta(hours=horas)
    g.tenant_db.execute(
        """INSERT INTO biblioteca_compartidos (tipo, entidad_id, token, expires_at, creado_por)
           VALUES ('documento', ?, ?, ?, ?)""",
        (doc_id, token, expires.isoformat(), session.get("user_id")),
    )
    g.tenant_db.commit()
    link = url_for("biblioteca.compartido_ver", token=token, _external=True)
    return jsonify(ok=True, link=link, expires=expires.isoformat())


# ═════════════════════════════════════════════════════════════════════════════
#  COMPARTIDO — Vista pública (sin login)
# ═════════════════════════════════════════════════════════════════════════════

@biblioteca_bp.route("/compartido/<token>")
def compartido_ver(token):
    """Public view for shared files/documents. No login required."""
    # We need to find the tenant from the token — check all tenant DBs
    # For simplicity, require tenant context
    tdb = getattr(g, "tenant_db", None)
    if not tdb:
        abort(404)
    share = tdb.execute(
        "SELECT * FROM biblioteca_compartidos WHERE token=? AND expires_at > datetime('now')",
        (token,),
    ).fetchone()
    if not share:
        return render_template("biblioteca/compartido_expired.html"), 404

    # Increment access count
    tdb.execute("UPDATE biblioteca_compartidos SET accesos=accesos+1 WHERE id=?", (share["id"],))
    tdb.commit()

    if share["tipo"] == "archivo":
        archivo = tdb.execute("SELECT * FROM biblioteca_archivos WHERE id=?", (share["entidad_id"],)).fetchone()
        if not archivo:
            abort(404)
        if boveda.es_de_boveda(archivo["archivo_url"]):
            return boveda.respuesta(archivo["archivo_url"], archivo["nombre_original"] or archivo["nombre"])
        from config import Config
        filepath = os.path.join(os.path.dirname(Config.UPLOAD_FOLDER), archivo["archivo_url"].lstrip("/"))
        if os.path.exists(filepath):
            return send_file(filepath, download_name=archivo["nombre_original"] or archivo["nombre"])
        abort(404)

    elif share["tipo"] == "documento":
        doc = tdb.execute("SELECT * FROM biblioteca_documentos WHERE id=?", (share["entidad_id"],)).fetchone()
        if not doc:
            abort(404)
        return render_template("biblioteca/compartido_doc.html", documento=doc, tenant=g.tenant)

    abort(404)
