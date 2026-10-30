"""
Habeas Corpus — modules/reportes.py
Reportes del despacho (pestaña «Reportes» dentro de Estadísticas): seis listas que se
revisan en pantalla y se descargan en Excel o PDF. Los filtros de uso frecuente se pueden
guardar con un nombre (y, si se quiere, ordenarlos en carpetas). Solo administradores.
"""
from __future__ import annotations

import io
import json
import re
import unicodedata
from datetime import datetime, timezone
from functools import wraps
from urllib.parse import quote
from xml.sax.saxutils import escape

from flask import (
    Blueprint, flash, g, jsonify, make_response, redirect, render_template, request,
    session, url_for,
)

from modules.asuntos import ABIERTOS, ESTADOS, ESTADOS_MENU, TIPOS_DERECHO, ZONA
from modules.auth import login_required, tenant_required

reportes_bp = Blueprint("reportes", __name__)

ADMINS = ("super_admin", "admin")

TIPOS_REPORTE = [
    {"id": "asuntos", "label": "Expedientes", "icon": "file-text",
     "desc": "Lista de expedientes con cliente, materia, estado, responsable y fechas."},
    {"id": "clientes", "label": "Clientes", "icon": "users",
     "desc": "Tus clientes con cuántos expedientes tienen y cuánto han pagado."},
    {"id": "estado_cuenta", "label": "Estado de cuenta", "icon": "credit-card",
     "desc": "Cada cobro por cliente: lo pagado, lo abonado y lo que falta por pagar."},
    {"id": "facturacion", "label": "Facturación", "icon": "dollar-sign",
     "desc": "Todo lo facturado en un periodo y cuánto ya se cobró (sin cancelados)."},
    {"id": "actividad", "label": "Actividad del equipo", "icon": "activity",
     "desc": "Qué hizo cada persona en el sistema y cuándo."},
    {"id": "vencimientos", "label": "Vencimientos", "icon": "alert-triangle",
     "desc": "Expedientes abiertos con fecha límite próxima o que ya se pasó."},
]
NOMBRES = {t["id"]: t["label"] for t in TIPOS_REPORTE}

ESTADOS_COBRO = {"pendiente": "Pendiente", "pagado": "Pagado", "vencido": "Vencido",
                 "abono": "Con abono", "cancelado": "Cancelado"}
TIPOS_CLIENTE = {"persona_fisica": "Persona física", "persona_moral": "Persona moral"}
ESTATUS_CLIENTE = {"activo": "Activo", "prospecto": "Prospecto", "inactivo": "Inactivo"}
ACCIONES = {"login": "Entró al sistema", "logout": "Salió del sistema", "crear": "Creó",
            "editar": "Editó", "eliminar": "Eliminó", "enviar_correo": "Envió un correo"}
ENTIDADES = {"asunto": "Expediente", "cobro": "Cobro", "biblioteca": "Documento",
             "correo": "Correo", "evento": "Agenda", "cliente": "Cliente"}

# Columnas con dinero: se muestran como $ y en Excel con formato de moneda
DINERO = {"total", "monto_abonado", "saldo", "total_cobrado"}


# ─────────────────────────────────────────────────────────────────────────────
#  Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _safe_int(v, default=None):
    try:
        return int(v) if v not in (None, "", "None") else default
    except (TypeError, ValueError):
        return default


def _fecha(v):
    """Una fecha AAAA-MM-DD válida o None (los filtros vienen del navegador)."""
    v = str(v or "")[:10]
    return v if re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) else None


def _solo_admin(f):
    """Reportes es solo para administradores. Las páginas regresan a Estadísticas."""
    @wraps(f)
    def decorated(*args, **kwargs):
        if session.get("role") not in ADMINS:
            if request.method == "GET":
                flash("Los reportes solo los ven los administradores del despacho.", "warning")
                return redirect(url_for("estadisticas.index"))
            return jsonify(ok=False, error="Solo los administradores pueden usar los reportes."), 403
        return f(*args, **kwargs)
    return decorated


def _get_abogados():
    try:
        return [dict(r) for r in g.global_db.execute(
            "SELECT id, nombre, apellidos FROM tenant_users "
            "WHERE tenant_id=? AND rol IN ('abogado','admin','super_admin') AND activo=1 ORDER BY nombre",
            (g.tenant["id"],),
        ).fetchall()]
    except Exception:
        return []


def _nombres_usuarios():
    try:
        return {r["id"]: f"{r['nombre']} {r['apellidos'] or ''}".strip() for r in g.global_db.execute(
            "SELECT id, nombre, apellidos FROM tenant_users WHERE tenant_id=?", (g.tenant["id"],)
        ).fetchall()}
    except Exception:
        return {}


def _etiqueta(mapa, valor):
    if not valor:
        return ""
    if valor in mapa:
        return mapa[valor]
    return str(valor).replace("_", " ").capitalize()


# ─────────────────────────────────────────────────────────────────────────────
#  Generadores: cada uno devuelve (filas, columnas, encabezados)
# ─────────────────────────────────────────────────────────────────────────────

def _gen_asuntos(tdb, filtros):
    q = """SELECT a.*, TRIM(c.nombre || ' ' || COALESCE(c.apellidos,'')) as cliente_nombre
           FROM asuntos a LEFT JOIN clientes c ON a.cliente_id=c.id WHERE 1=1"""
    p = []
    if filtros.get("estado") in ESTADOS:
        q += " AND a.estado=?"; p.append(filtros["estado"])
    if filtros.get("tipo"):
        q += " AND a.tipo_asunto=?"; p.append(str(filtros["tipo"]))
    if _safe_int(filtros.get("abogado")):
        q += " AND a.abogado_responsable=?"; p.append(_safe_int(filtros["abogado"]))
    if _fecha(filtros.get("desde")):
        q += " AND a.fecha_alta >= ?"; p.append(_fecha(filtros["desde"]))
    if _fecha(filtros.get("hasta")):
        q += " AND a.fecha_alta <= ? || ' 23:59:59'"; p.append(_fecha(filtros["hasta"]))
    q += " ORDER BY a.fecha_alta DESC LIMIT 500"
    rows = [dict(r) for r in tdb.execute(q, p).fetchall()]
    nombres = _nombres_usuarios()
    for r in rows:
        r["estado"] = ESTADOS.get(r.get("estado"), {}).get("label") or _etiqueta({}, r.get("estado"))
        r["tipo_asunto"] = TIPOS_DERECHO.get(r.get("tipo_asunto"), {}).get("label") or _etiqueta({}, r.get("tipo_asunto"))
        r["abogado"] = nombres.get(r.get("abogado_responsable"), "")
        r["fecha_alta"] = (r.get("fecha_alta") or "")[:10]
    columns = ["numero_expediente", "titulo", "cliente_nombre", "tipo_asunto", "estado",
               "abogado", "fecha_alta", "fecha_limite"]
    headers = ["Expediente", "Título", "Cliente", "Materia", "Estado", "Responsable",
               "Alta", "Fecha límite"]
    return rows, columns, headers


def _gen_clientes(tdb, filtros):
    q = """SELECT c.*,
           (SELECT COUNT(*) FROM asuntos WHERE cliente_id=c.id) as total_asuntos,
           (SELECT COALESCE(SUM(COALESCE(total,monto)),0) FROM cobros WHERE cliente_id=c.id AND estado='pagado') as total_cobrado
           FROM clientes c WHERE 1=1"""
    p = []
    if filtros.get("tipo") in TIPOS_CLIENTE:
        q += " AND c.tipo=?"; p.append(filtros["tipo"])
    if filtros.get("estatus") in ESTATUS_CLIENTE:
        q += " AND c.estatus=?"; p.append(filtros["estatus"])
    if _fecha(filtros.get("desde")):
        q += " AND c.fecha_alta >= ?"; p.append(_fecha(filtros["desde"]))
    if _fecha(filtros.get("hasta")):
        q += " AND c.fecha_alta <= ? || ' 23:59:59'"; p.append(_fecha(filtros["hasta"]))
    q += " ORDER BY c.nombre, c.apellidos LIMIT 500"
    rows = [dict(r) for r in tdb.execute(q, p).fetchall()]
    for r in rows:
        r["nombre_completo"] = (r.get("razon_social") if r.get("tipo") == "persona_moral" and r.get("razon_social")
                                else f"{r.get('nombre') or ''} {r.get('apellidos') or ''}".strip())
        r["tipo"] = _etiqueta(TIPOS_CLIENTE, r.get("tipo"))
        r["estatus"] = _etiqueta(ESTATUS_CLIENTE, r.get("estatus"))
        r["telefono"] = r.get("celular") or r.get("telefono") or ""
        r["fecha_alta"] = (r.get("fecha_alta") or "")[:10]
    columns = ["nombre_completo", "tipo", "estatus", "email", "telefono", "total_asuntos",
               "total_cobrado", "fecha_alta"]
    headers = ["Cliente", "Tipo", "Estatus", "Correo", "Teléfono", "Expedientes",
               "Total pagado", "Alta"]
    return rows, columns, headers


def _gen_estado_cuenta(tdb, filtros):
    q = """SELECT co.*, TRIM(c.nombre || ' ' || COALESCE(c.apellidos,'')) as cliente_nombre,
           a.titulo as asunto_titulo
           FROM cobros co
           LEFT JOIN clientes c ON co.cliente_id=c.id
           LEFT JOIN asuntos a ON co.asunto_id=a.id WHERE 1=1"""
    p = []
    if _safe_int(filtros.get("cliente_id")):
        q += " AND co.cliente_id=?"; p.append(_safe_int(filtros["cliente_id"]))
    if filtros.get("estado") in ESTADOS_COBRO:
        q += " AND co.estado=?"; p.append(filtros["estado"])
    if _fecha(filtros.get("desde")):
        q += " AND co.fecha_emision >= ?"; p.append(_fecha(filtros["desde"]))
    if _fecha(filtros.get("hasta")):
        q += " AND co.fecha_emision <= ?"; p.append(_fecha(filtros["hasta"]))
    q += " ORDER BY cliente_nombre, co.fecha_emision DESC LIMIT 500"
    rows = [dict(r) for r in tdb.execute(q, p).fetchall()]
    for r in rows:
        total = float(r.get("total") or r.get("monto") or 0)
        abonado = float(r.get("monto_abonado") or 0)
        r["total"] = total
        r["monto_abonado"] = total if r.get("estado") == "pagado" else abonado
        r["saldo"] = 0.0 if r.get("estado") in ("pagado", "cancelado") else max(total - abonado, 0.0)
        r["folio"] = r.get("numero_factura") or f"#{r['id']}"
        r["estado"] = _etiqueta(ESTADOS_COBRO, r.get("estado"))
    columns = ["cliente_nombre", "asunto_titulo", "folio", "concepto", "total", "monto_abonado",
               "saldo", "estado", "fecha_emision", "fecha_vencimiento"]
    headers = ["Cliente", "Expediente", "Folio", "Concepto", "Total", "Pagado",
               "Por pagar", "Estado", "Emisión", "Vence"]
    return rows, columns, headers


def _gen_facturacion(tdb, filtros):
    q = """SELECT co.*, TRIM(c.nombre || ' ' || COALESCE(c.apellidos,'')) as cliente_nombre
           FROM cobros co
           LEFT JOIN clientes c ON co.cliente_id=c.id
           WHERE co.estado != 'cancelado'"""
    p = []
    if _fecha(filtros.get("desde")):
        q += " AND co.fecha_emision >= ?"; p.append(_fecha(filtros["desde"]))
    if _fecha(filtros.get("hasta")):
        q += " AND co.fecha_emision <= ?"; p.append(_fecha(filtros["hasta"]))
    if filtros.get("moneda") in ("MXN", "USD"):
        q += " AND co.moneda=?"; p.append(filtros["moneda"])
    q += " ORDER BY co.fecha_emision DESC LIMIT 500"
    rows = [dict(r) for r in tdb.execute(q, p).fetchall()]
    for r in rows:
        total = float(r.get("total") or r.get("monto") or 0)
        r["total"] = total
        r["monto_abonado"] = total if r.get("estado") == "pagado" else float(r.get("monto_abonado") or 0)
        r["folio"] = r.get("numero_factura") or f"#{r['id']}"
        r["estado"] = _etiqueta(ESTADOS_COBRO, r.get("estado"))
    columns = ["folio", "cliente_nombre", "concepto", "moneda", "total", "monto_abonado",
               "estado", "fecha_emision", "fecha_pago"]
    headers = ["Folio", "Cliente", "Concepto", "Moneda", "Total", "Cobrado", "Estado",
               "Emisión", "Pago"]
    return rows, columns, headers


def _gen_actividad(tdb, filtros):
    q = """SELECT al.*, a.titulo as asunto_titulo
           FROM actividad_log al
           LEFT JOIN asuntos a ON al.entidad='asunto' AND al.entidad_id=a.id
           WHERE 1=1"""
    p = []
    if _safe_int(filtros.get("abogado")):
        q += " AND al.user_id=?"; p.append(_safe_int(filtros["abogado"]))
    # La bitácora guarda la hora en UTC; los filtros y la pantalla van en hora de la Ciudad de México
    if _fecha(filtros.get("desde")):
        q += " AND datetime(al.fecha, '-6 hours') >= ?"; p.append(_fecha(filtros["desde"]))
    if _fecha(filtros.get("hasta")):
        q += " AND datetime(al.fecha, '-6 hours') <= ? || ' 23:59:59'"; p.append(_fecha(filtros["hasta"]))
    if filtros.get("accion"):
        q += " AND al.accion=?"; p.append(str(filtros["accion"]))
    q += " ORDER BY al.fecha DESC LIMIT 500"
    rows = [dict(r) for r in tdb.execute(q, p).fetchall()]
    nombres = _nombres_usuarios()
    for r in rows:
        r["persona"] = nombres.get(r.get("user_id"), "—")
        r["accion"] = _etiqueta(ACCIONES, r.get("accion"))
        r["entidad"] = _etiqueta(ENTIDADES, r.get("entidad"))
        r["fecha"] = _hora_local(r.get("fecha"))
    columns = ["fecha", "persona", "accion", "entidad", "asunto_titulo"]
    headers = ["Fecha y hora", "Persona", "Qué hizo", "Sobre", "Expediente"]
    return rows, columns, headers


def _gen_vencimientos(tdb, filtros):
    marcas = ",".join("?" * len(ABIERTOS))
    q = f"""SELECT a.*, TRIM(c.nombre || ' ' || COALESCE(c.apellidos,'')) as cliente_nombre,
           CAST(julianday(a.fecha_limite) - julianday(date('now','localtime')) AS INTEGER) as dias
           FROM asuntos a LEFT JOIN clientes c ON a.cliente_id=c.id
           WHERE a.estado IN ({marcas}) AND a.fecha_limite IS NOT NULL AND a.fecha_limite != ''"""
    p = list(ABIERTOS)
    rango = filtros.get("rango") or "30"
    hoy = "date('now','localtime')"
    if rango == "vencidos":
        q += f" AND a.fecha_limite < {hoy}"
    elif rango in ("7", "30", "90"):
        q += f" AND a.fecha_limite BETWEEN {hoy} AND date('now','localtime', '+{int(rango)} days')"
    if _safe_int(filtros.get("abogado")):
        q += " AND a.abogado_responsable=?"; p.append(_safe_int(filtros["abogado"]))
    q += " ORDER BY a.fecha_limite ASC LIMIT 500"
    rows = [dict(r) for r in tdb.execute(q, p).fetchall()]
    nombres = _nombres_usuarios()
    for r in rows:
        d = r.get("dias")
        r["faltan"] = ("" if d is None else "Hoy" if d == 0 else
                       f"Hace {-d} día{'s' if d != -1 else ''}" if d < 0 else f"{d} día{'s' if d != 1 else ''}")
        r["abogado"] = nombres.get(r.get("abogado_responsable"), "")
        r["tipo_asunto"] = TIPOS_DERECHO.get(r.get("tipo_asunto"), {}).get("label") or _etiqueta({}, r.get("tipo_asunto"))
    columns = ["fecha_limite", "faltan", "numero_expediente", "titulo", "cliente_nombre",
               "tipo_asunto", "abogado"]
    headers = ["Fecha límite", "Faltan", "Expediente", "Título", "Cliente", "Materia", "Responsable"]
    return rows, columns, headers


GENERATORS = {
    "asuntos": _gen_asuntos,
    "clientes": _gen_clientes,
    "estado_cuenta": _gen_estado_cuenta,
    "facturacion": _gen_facturacion,
    "actividad": _gen_actividad,
    "vencimientos": _gen_vencimientos,
}


def _peticion():
    """(tipo, filtros, nombre) desde JSON (pantalla) o desde un formulario (descargas)."""
    data = request.get_json(silent=True)
    if data is None:
        data = request.form.to_dict()
        try:
            data["filtros"] = json.loads(data.get("filtros") or "{}")
        except ValueError:
            data["filtros"] = {}
    filtros = data.get("filtros") if isinstance(data.get("filtros"), dict) else {}
    tipo = data.get("tipo", "")
    nombre = (str(data.get("nombre") or "").strip() or NOMBRES.get(tipo, "Reporte"))[:80]
    return tipo, filtros, nombre


def _archivo(nombre, ext):
    """Nombre de archivo seguro + versión con acentos para el navegador."""
    base = unicodedata.normalize("NFKD", nombre).encode("ascii", "ignore").decode()
    base = re.sub(r"[^A-Za-z0-9 _-]+", "", base).strip().replace(" ", "_") or "reporte"
    fecha = datetime.now(ZONA).strftime("%Y-%m-%d")
    completo = f"{nombre} {fecha}.{ext}"
    return f'attachment; filename="{base}_{fecha}.{ext}"; filename*=UTF-8\'\'{quote(completo)}'


def _hora_local(v):
    try:
        utc = datetime.fromisoformat(str(v)[:19]).replace(tzinfo=timezone.utc)
        return utc.astimezone(ZONA).strftime("%Y-%m-%d %H:%M")
    except (TypeError, ValueError):
        return str(v or "")[:16]


def _texto(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return f"${v:,.2f}"
    return str(v)


# ─────────────────────────────────────────────────────────────────────────────
#  Rutas
# ─────────────────────────────────────────────────────────────────────────────

@reportes_bp.route("/")
@login_required
@tenant_required
@_solo_admin
def index():
    tdb = g.tenant_db
    carpetas = [dict(r) for r in tdb.execute("SELECT * FROM reportes_carpetas ORDER BY nombre").fetchall()]
    guardados = [dict(r) for r in tdb.execute(
        "SELECT * FROM reportes_guardados ORDER BY fecha_mod DESC"
    ).fetchall()]
    for r in guardados:
        r["tipo_label"] = NOMBRES.get(r["tipo"], r["tipo"])
        try:
            r["filtros"] = json.loads(r.get("filtros_json") or "{}")
        except ValueError:
            r["filtros"] = {}
    clientes = [dict(r) for r in tdb.execute(
        "SELECT id, TRIM(nombre || ' ' || COALESCE(apellidos,'')) AS nombre FROM clientes ORDER BY nombre, apellidos"
    ).fetchall()]
    acciones = [r[0] for r in tdb.execute(
        "SELECT DISTINCT accion FROM actividad_log WHERE accion IS NOT NULL AND accion != '' ORDER BY accion"
    ).fetchall()]

    return render_template(
        "reportes/index.html",
        tipos_reporte=TIPOS_REPORTE,
        materias=[(k, v["label"]) for k, v in TIPOS_DERECHO.items()],
        estados_asunto=[(k, ESTADOS[k]["label"]) for k in ESTADOS_MENU],
        estados_cobro=ESTADOS_COBRO,
        tipos_cliente=TIPOS_CLIENTE,
        estatus_cliente=ESTATUS_CLIENTE,
        acciones=[(a, _etiqueta(ACCIONES, a)) for a in acciones],
        clientes=clientes,
        carpetas=carpetas,
        guardados=guardados,
        abogados=_get_abogados(),
    )


# ── Carpetas ──

@reportes_bp.route("/carpeta/crear", methods=["POST"])
@login_required
@tenant_required
@_solo_admin
def carpeta_crear():
    data = request.get_json(silent=True) or request.form
    nombre = (data.get("nombre") or "").strip()[:80]
    if not nombre:
        return jsonify(ok=False, error="Escribe el nombre de la carpeta."), 400
    cur = g.tenant_db.execute(
        "INSERT INTO reportes_carpetas (nombre, creado_por) VALUES (?,?)",
        (nombre, session.get("user_id")),
    )
    g.tenant_db.commit()
    return jsonify(ok=True, id=cur.lastrowid)


@reportes_bp.route("/carpeta/<int:cid>/eliminar", methods=["POST"])
@login_required
@tenant_required
@_solo_admin
def carpeta_eliminar(cid):
    g.tenant_db.execute("UPDATE reportes_guardados SET carpeta_id=NULL WHERE carpeta_id=?", (cid,))
    g.tenant_db.execute("DELETE FROM reportes_carpetas WHERE id=?", (cid,))
    g.tenant_db.commit()
    return jsonify(ok=True)


# ── Reportes guardados ──

@reportes_bp.route("/guardar", methods=["POST"])
@login_required
@tenant_required
@_solo_admin
def guardar():
    data = request.get_json(silent=True) or request.form
    nombre = (data.get("nombre") or "").strip()[:120]
    tipo = data.get("tipo", "")
    filtros = data.get("filtros") if isinstance(data.get("filtros"), dict) else {}
    carpeta_id = _safe_int(data.get("carpeta_id"))
    rid = _safe_int(data.get("id"))

    if not nombre:
        return jsonify(ok=False, error="Ponle un nombre al reporte."), 400
    if tipo not in GENERATORS:
        return jsonify(ok=False, error="Elige primero un tipo de reporte."), 400

    tdb = g.tenant_db
    if carpeta_id and not tdb.execute("SELECT 1 FROM reportes_carpetas WHERE id=?", (carpeta_id,)).fetchone():
        carpeta_id = None
    filtros_json = json.dumps(filtros, ensure_ascii=False)
    if rid and tdb.execute("SELECT 1 FROM reportes_guardados WHERE id=?", (rid,)).fetchone():
        tdb.execute(
            "UPDATE reportes_guardados SET nombre=?, tipo=?, filtros_json=?, carpeta_id=?, fecha_mod=CURRENT_TIMESTAMP WHERE id=?",
            (nombre, tipo, filtros_json, carpeta_id, rid),
        )
    else:
        rid = tdb.execute(
            "INSERT INTO reportes_guardados (nombre, tipo, filtros_json, carpeta_id, creado_por) VALUES (?,?,?,?,?)",
            (nombre, tipo, filtros_json, carpeta_id, session.get("user_id")),
        ).lastrowid
    tdb.commit()
    return jsonify(ok=True, id=rid)


@reportes_bp.route("/guardado/<int:rid>/eliminar", methods=["POST"])
@login_required
@tenant_required
@_solo_admin
def guardado_eliminar(rid):
    g.tenant_db.execute("DELETE FROM reportes_guardados WHERE id=?", (rid,))
    g.tenant_db.commit()
    return jsonify(ok=True)


# ── Ver en pantalla ──

@reportes_bp.route("/generar", methods=["POST"])
@login_required
@tenant_required
@_solo_admin
def generar():
    tipo, filtros, _ = _peticion()
    gen = GENERATORS.get(tipo)
    if not gen:
        return jsonify(ok=False, error="Elige un tipo de reporte."), 400
    rows, columns, headers = gen(g.tenant_db, filtros)
    rows = [{c: r.get(c) for c in columns} for r in rows]
    return jsonify(ok=True, rows=rows, columns=columns, headers=headers, total=len(rows),
                   dinero=[c for c in columns if c in DINERO])


# ── Descargar en Excel ──

@reportes_bp.route("/exportar/excel", methods=["POST"])
@login_required
@tenant_required
@_solo_admin
def exportar_excel():
    tipo, filtros, nombre = _peticion()
    gen = GENERATORS.get(tipo)
    if not gen:
        return jsonify(ok=False, error="Elige un tipo de reporte."), 400
    rows, columns, headers = gen(g.tenant_db, filtros)

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "", nombre)[:31] or "Reporte"
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="142E59")
        cell.alignment = Alignment(vertical="center")
    for ri, row in enumerate(rows, 2):
        for ci, col in enumerate(columns, 1):
            v = row.get(col)
            cell = ws.cell(row=ri, column=ci, value=v if v not in (None, "") else None)
            if col in DINERO:
                cell.number_format = '"$"#,##0.00'
    ws.freeze_panes = "A2"
    for ci, col_cells in enumerate(ws.columns, 1):
        largo = max(len(str(c.value or "")) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(largo + 2, 10), 50)

    buf = io.BytesIO()
    wb.save(buf)
    resp = make_response(buf.getvalue())
    resp.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    resp.headers["Content-Disposition"] = _archivo(nombre, "xlsx")
    return resp


# ── Descargar en PDF ──

@reportes_bp.route("/exportar/pdf", methods=["POST"])
@login_required
@tenant_required
@_solo_admin
def exportar_pdf():
    tipo, filtros, nombre = _peticion()
    gen = GENERATORS.get(tipo)
    if not gen:
        return jsonify(ok=False, error="Elige un tipo de reporte."), 400
    rows, columns, headers = gen(g.tenant_db, filtros)

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    marino = colors.HexColor("#142E59")
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), leftMargin=1.4 * cm, rightMargin=1.4 * cm,
                            topMargin=1.4 * cm, bottomMargin=1.4 * cm, title=nombre)
    st_tit = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=16, textColor=marino, leading=20)
    st_meta = ParagraphStyle("m", fontName="Helvetica", fontSize=9, textColor=colors.HexColor("#60718D"), leading=12)
    st_th = ParagraphStyle("th", fontName="Helvetica-Bold", fontSize=8, textColor=colors.white, leading=10)
    st_td = ParagraphStyle("td", fontName="Helvetica", fontSize=8, textColor=marino, leading=10)

    despacho = g.tenant["nombre_despacho"] or "Habeas Corpus"
    total = len(rows)
    partes = [Paragraph(escape(nombre), st_tit),
              Paragraph(escape(f"{despacho} · Generado el {datetime.now(ZONA).strftime('%d/%m/%Y %H:%M')} · "
                               f"{total} registro{'s' if total != 1 else ''}"), st_meta),
              Spacer(1, 10)]
    datos = [[Paragraph(escape(h), st_th) for h in headers]]
    for r in rows:
        datos.append([Paragraph(escape(_texto(r.get(c))), st_td) for c in columns])
    if not rows:
        datos.append([Paragraph("Sin resultados con estos filtros.", st_td)] + [""] * (len(headers) - 1))
    ancho = doc.width / len(headers)
    tabla = Table(datos, colWidths=[ancho] * len(headers), repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), marino),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, colors.HexColor("#E5E2DB")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAF8F4")]),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    if not rows:
        estilo.append(("SPAN", (0, 1), (-1, 1)))
    tabla.setStyle(TableStyle(estilo))
    partes.append(tabla)
    doc.build(partes)

    resp = make_response(buf.getvalue())
    resp.headers["Content-Type"] = "application/pdf"
    resp.headers["Content-Disposition"] = _archivo(nombre, "pdf")
    return resp
