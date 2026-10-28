"""Extrae texto de los documentos del despacho, con OCR local para escaneos."""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field

MIN_TEXTO_PAGINA = 40          # menos caracteres que esto en una página de PDF = probablemente escaneada
ESCALA_OCR = 3                 # ≈216 dpi al rasterizar páginas escaneadas
IMAGENES = {"png", "jpg", "jpeg", "gif", "webp", "tif", "tiff"}
SOPORTADOS = {"pdf", "docx", "xlsx", "txt", "csv", "xml"} | IMAGENES


class FaltaOCR(Exception):
    """El documento necesita OCR y Tesseract no está instalado."""


@dataclass
class Resultado:
    paginas: list[tuple[int, str]] = field(default_factory=list)   # (número de página, texto)
    uso_ocr: bool = False

    @property
    def caracteres(self) -> int:
        return sum(len(t) for _, t in self.paginas)


def ocr_disponible() -> bool:
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _ocr(imagen) -> str:
    import pytesseract
    if not ocr_disponible():
        raise FaltaOCR()
    idiomas = "spa" if "spa" in pytesseract.get_languages(config="") else "eng"
    return pytesseract.image_to_string(imagen, lang=idiomas)


def _pdf(datos: bytes, res: Resultado) -> None:
    from pypdf import PdfReader
    lector = PdfReader(io.BytesIO(datos))
    if lector.is_encrypted:
        lector.decrypt("")
    pendientes_ocr = []
    for i, pagina in enumerate(lector.pages, start=1):
        try:
            texto = pagina.extract_text() or ""
        except Exception:
            texto = ""
        if len(texto.strip()) >= MIN_TEXTO_PAGINA:
            res.paginas.append((i, texto))
        else:
            pendientes_ocr.append(i)
    if not pendientes_ocr:
        return
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(datos)
    try:
        for i in pendientes_ocr:
            imagen = doc[i - 1].render(scale=ESCALA_OCR).to_pil()
            res.paginas.append((i, _ocr(imagen)))
            res.uso_ocr = True
    finally:
        doc.close()
    res.paginas.sort()


def _imagen(datos: bytes, res: Resultado) -> None:
    from PIL import Image, ImageSequence
    with Image.open(io.BytesIO(datos)) as img:
        for i, cuadro in enumerate(ImageSequence.Iterator(img), start=1):
            res.paginas.append((i, _ocr(cuadro.convert("RGB"))))
    res.uso_ocr = True


def _docx(datos: bytes, res: Resultado) -> None:
    from docx import Document
    doc = Document(io.BytesIO(datos))
    partes = [p.text for p in doc.paragraphs if p.text.strip()]
    for tabla in doc.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
            if celdas:
                partes.append(" | ".join(celdas))
    res.paginas.append((1, "\n".join(partes)))


def _xlsx(datos: bytes, res: Resultado) -> None:
    from openpyxl import load_workbook
    libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    for i, hoja in enumerate(libro.worksheets, start=1):
        filas = []
        for fila in hoja.iter_rows(values_only=True):
            valores = [str(v) for v in fila if v not in (None, "")]
            if valores:
                filas.append(" | ".join(valores))
        if filas:
            res.paginas.append((i, f"Hoja {hoja.title}\n" + "\n".join(filas)))


def _texto_plano(datos: bytes) -> str:
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            return datos.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return ""


def extraer(datos: bytes, ext: str) -> Resultado:
    """Devuelve el texto por página. Lanza FaltaOCR si hace falta OCR y no está instalado."""
    ext = ext.lower()
    res = Resultado()
    if ext == "pdf":
        _pdf(datos, res)
    elif ext in IMAGENES:
        _imagen(datos, res)
    elif ext == "docx":
        _docx(datos, res)
    elif ext == "xlsx":
        _xlsx(datos, res)
    elif ext == "csv":
        filas = csv.reader(io.StringIO(_texto_plano(datos)))
        res.paginas.append((1, "\n".join(" | ".join(f) for f in filas)))
    elif ext in ("txt", "xml"):
        res.paginas.append((1, _texto_plano(datos)))
    return res
